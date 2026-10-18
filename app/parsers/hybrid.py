"""XLSX 受控视觉增强：确定性 Native 主结果、单个 Bundle 及原 Anchor 合并。"""

import hashlib
import json
import posixpath
import re
import struct
from dataclasses import dataclass, field
from collections.abc import Callable
from pathlib import Path
from typing import Literal
from uuid import UUID, uuid5
from zipfile import BadZipFile, ZipFile

import pymupdf
from lxml import etree
from openpyxl.utils.cell import range_boundaries
from openpyxl.formula.translate import Translator
from pydantic import BaseModel, ConfigDict, Field, field_validator

from app.parsers.common import NoTextError, ParserError
from app.parsers.routing.inspector import xml_part
from app.parsers.xlsx import XlsxParser
from app.schemas.parsed_document import ParsedBlock, ParsedDocument

VERSION = "xlsx-hybrid-v1"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
IMAGE = "http://schemas.openxmlformats.org/drawingml/2006/main"
CHART = "http://schemas.openxmlformats.org/drawingml/2006/chart"


class HybridPreflightError(ParserError):
    """固定安全拒绝原因；只有已知能力缺口允许在提交前改走整份云端。"""

    WHOLE_COMPATIBLE = frozenset({
        "HYBRID_UNSUPPORTED_STRUCTURE", "HYBRID_UNSUPPORTED_ANCHOR",
        "HYBRID_UNSUPPORTED_CHART", "HYBRID_UNSUPPORTED_IMAGE_TRANSFORM",
        "HYBRID_UNSUPPORTED_IMAGE_FORMAT", "HYBRID_UNSUPPORTED_DRAWING",
        "HYBRID_UNSUPPORTED_SHEET",
    })

    def __init__(self, code: str = "HYBRID_UNSUPPORTED_STRUCTURE") -> None:
        """仅接收安全固定错误码，不泄露路径、XML 或文档正文。"""

        allowed = self.WHOLE_COMPATIBLE | {
            "HYBRID_RESOURCE_LIMIT", "HYBRID_EXTERNAL_ASSET_MISSING",
            "HYBRID_INVALID_PACKAGE", "HYBRID_INVALID_IMAGE", "HYBRID_ORPHAN_PART",
        }
        self.code = code if code in allowed else "HYBRID_INVALID_PACKAGE"
        super().__init__(f"ROUTER_FAILED[{self.code}]：Hybrid 预检或完整性校验未通过")

    @property
    def whole_compatible(self) -> bool:
        """只暴露受控能力缺口；损坏、缺附件、孤立部件和资源超限不能转付费云请求。"""

        return self.code in self.WHOLE_COMPATIBLE


class HybridPolicy(BaseModel):
    """检查点冻结的资源阈值；重试不能因环境变量改变而漏掉原候选。"""

    model_config = ConfigDict(extra="forbid", frozen=True, strict=True)
    max_candidates: int = Field(default=64, ge=1, le=256)
    max_pages: int = Field(default=32, ge=1, le=200)
    min_side: int = Field(default=32, ge=1, le=4096)
    max_pixels: int = Field(default=16000000, ge=1, le=40000000)
    max_image_bytes: int = Field(default=16 * 1024**2, ge=1, le=64 * 1024**2)
    max_bundle_bytes: int = Field(default=50 * 1024**2, ge=1, le=200 * 1024**2)
    xml_limit: int = Field(default=16 * 1024**2, ge=1, le=64 * 1024**2)


class VisualEntry(BaseModel):
    """来源账本：Anchor 为零基行列及偏移，不以临时 PDF 页冒充 Excel 页。"""

    model_config = ConfigDict(extra="forbid")
    visual_id: UUID
    source_doc_id: UUID
    source_type: Literal["embedded_image", "chart"]
    source_section: str
    sheet_index: int = Field(ge=1)
    source_anchor: tuple[int, ...]
    position: int = Field(ge=1)
    content_hash: str = Field(pattern=r"^[0-9a-f]{64}$")
    media_type: str
    width: int = Field(default=0, ge=0)
    height: int = Field(default=0, ge=0)
    bundle_page: int | None = Field(default=None, ge=1, le=200)
    state: Literal["pending", "parsed", "filtered", "reused", "failed"]
    filter_reason: Literal["small", "solid"] | None = None
    reused_from: UUID | None = None
    block_ids: list[UUID] = Field(default_factory=list)
    cache_hit: bool = False
    cache_result_hash: str | None = Field(default=None, pattern=r"^[0-9a-f]{64}$")

    @field_validator("source_anchor")
    @classmethod
    def validate_anchor(cls, value: tuple[int, ...]) -> tuple[int, ...]:
        """拒绝缺项或负数 Anchor，保持包内单/双单元格位置可解释。"""

        if len(value) not in {4, 8} or any(component < 0 for component in value):
            raise ValueError("Hybrid Anchor 必须为单/双单元格非负坐标")
        return value


class BundleManifest(BaseModel):
    """私有有界映射；不进入日志、Celery 消息、API 请求或公开响应。"""

    model_config = ConfigDict(extra="forbid")
    version: Literal["xlsx-hybrid-v1"] = VERSION
    policy: HybridPolicy
    sheet_names: list[str] = Field(max_length=1000)
    entries: list[VisualEntry] = Field(max_length=256)
    bundle_page_count: int = Field(ge=0, le=200)
    bundle_hash: str = Field(pattern=r"^[0-9a-f]{64}$")
    native_hash: str = Field(pattern=r"^[0-9a-f]{64}$")


@dataclass(repr=False)
class HybridPlan:
    """仅本次处理内存暂存的主结果、图表及 Bundle，不是 Parse Cache。"""

    native: ParsedDocument
    manifest: BundleManifest
    bundle: bytes = field(repr=False)
    charts: dict[UUID, ParsedBlock] = field(repr=False)
    cached: dict[str, list[ParsedBlock]] = field(default_factory=dict, repr=False)


def digest(value: bytes) -> str:
    """计算精确内容 Hash，不用 OCR 文本判断图片相同。"""

    return hashlib.sha256(value).hexdigest()


def relationships(archive: ZipFile, owner: str, limit: int) -> dict[str, str]:
    """解析包内关系；超链接不当附件抓取，视觉外链/缺部件/越界路径安全失败。"""

    parent, name = posixpath.split(owner)
    rel_name = posixpath.join(parent, "_rels", name + ".rels")
    if rel_name not in archive.namelist():
        return {}
    tree = xml_part(archive, rel_name, limit)
    result = {}
    for node in tree:
        target, identifier = node.get("Target", ""), node.get("Id", "")
        if node.get("TargetMode", "Internal") != "Internal":
            if node.get("Type") == R + "/hyperlink":
                continue
            raise HybridPreflightError("HYBRID_EXTERNAL_ASSET_MISSING")
        if (
            not target or "\\" in target or ":" in target or "?" in target
            or "#" in target or not identifier or identifier in result
        ):
            raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
        part = posixpath.normpath(target.lstrip("/") if target.startswith("/")
                                 else posixpath.join(parent, target))
        if not part.startswith("xl/") or part not in archive.namelist():
            raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
        result[identifier] = part
    return result


def image_dimensions(content: bytes) -> tuple[int, int, str]:
    """解码前检查 PNG/JPEG 头；已知其他图像格式可 Whole，畸形字节直接失败。"""

    if content.startswith(b"\x89PNG\r\n\x1a\n") and len(content) >= 24:
        if content[12:16] != b"IHDR":
            raise HybridPreflightError("HYBRID_INVALID_IMAGE")
        width, height = struct.unpack(">II", content[16:24])
        return width, height, "image/png"
    if content.startswith(b"\xff\xd8"):
        offset = 2
        while offset + 4 <= len(content):
            if content[offset] != 255:
                raise HybridPreflightError("HYBRID_INVALID_IMAGE")
            while offset < len(content) and content[offset] == 255:
                offset += 1
            if offset >= len(content):
                break
            marker = content[offset]
            offset += 1
            if marker in {0xD9, 0xDA}:
                break
            if marker == 0x01 or 0xD0 <= marker <= 0xD7:
                continue
            length = int.from_bytes(content[offset:offset + 2], "big")
            if length < 2 or offset + length > len(content):
                raise HybridPreflightError("HYBRID_INVALID_IMAGE")
            if marker in {0xC0, 0xC1, 0xC2} and length >= 8:
                height, width = struct.unpack(">HH", content[offset + 3:offset + 7])
                return width, height, "image/jpeg"
            offset += length
    if (content.startswith((b"GIF87a", b"GIF89a", b"BM", b"II*\x00", b"MM\x00*"))
        or (content.startswith(b"RIFF") and content[8:12] == b"WEBP")):
        raise HybridPreflightError("HYBRID_UNSUPPORTED_IMAGE_FORMAT")
    raise HybridPreflightError("HYBRID_INVALID_IMAGE")


def anchor_coordinates(node) -> tuple[int, ...]:
    """保留单/双单元格锚点；绝对位置或未知变换不猜测 Excel Anchor。"""

    if node.tag not in {f"{{{XDR}}}oneCellAnchor", f"{{{XDR}}}twoCellAnchor"}:
        if node.tag == f"{{{XDR}}}absoluteAnchor":
            raise HybridPreflightError("HYBRID_UNSUPPORTED_ANCHOR")
        raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
    values = []
    for marker in ("from", "to"):
        item = node.find(f"{{{XDR}}}{marker}")
        if item is None:
            if marker == "to" and node.tag.endswith("oneCellAnchor"):
                break
            raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
        for tag in ("row", "col", "rowOff", "colOff"):
            raw = item.findtext(f"{{{XDR}}}{tag}")
            if raw is None or not re.fullmatch(r"\d{1,10}", raw):
                raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
            values.append(int(raw))
        if values[-4] > 1048575 or values[-3] > 16383:
            raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
    if len(values) == 8 and (values[4] < values[0] or values[5] < values[1]):
        raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
    return tuple(values)


def chart_text(tree, cell_values: dict[str, dict[str, str]]) -> str:
    """提取基础 Chart 引用与原式并保留常见线条/颜色样式；复杂视觉语义交整份云端。"""

    tags = {etree.QName(node).localname for node in tree.iter()}
    unsupported = {
        "extLst", "userShapes", "externalData", "pivotSource", "trendline", "errBars",
        "multiLvlStrRef", "dLbls", "dTable", "pictureOptions", "scatterChart",
        "bubbleChart", "surfaceChart", "stockChart", "radarChart", "dispUnits", "logBase", "dPt",
    }
    if tags & unsupported or any(len(n) for n in tree.findall(".//{*}layout")):
        raise HybridPreflightError("HYBRID_UNSUPPORTED_CHART")
    style_lines = []
    safe_style = {"spPr", "ln", "solidFill", "noFill", "srgbClr", "schemeClr",
                  "prstDash", "lumMod", "lumOff", "alpha", "round", "bevel", "miter"}
    for style in tree.findall(".//{*}spPr"):
        if any(etree.QName(n).localname not in safe_style for n in style.iter()):
            raise HybridPreflightError("HYBRID_UNSUPPORTED_CHART")
        if len(style):
            style_lines.append("样式：" + json.dumps([
                {"type": etree.QName(n).localname, "attributes": dict(sorted(n.attrib.items()))}
                for n in style.iter()
            ], ensure_ascii=False, sort_keys=True))
    plot = tree.find(".//{*}plotArea")
    if plot is None:
        raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
    charts = [node for node in plot if etree.QName(node).localname.endswith("Chart")]
    if len(charts) != 1 or etree.QName(charts[0]).localname not in {
        "barChart", "lineChart", "pieChart", "doughnutChart", "areaChart"
    }:
        raise HybridPreflightError("HYBRID_UNSUPPORTED_CHART")
    lines = ["图表类型：" + etree.QName(charts[0]).localname] + style_lines
    for title in tree.findall(".//{*}title"):
        if title.find(".//{*}f") is not None:
            raise HybridPreflightError("HYBRID_UNSUPPORTED_CHART")
        title_text = "".join(node.text or "" for node in title.findall(".//{*}t"))
        if title_text:
            lines.append("标题：" + title_text)
    # 把确定性分组方向/数字格式保留下来，不把原始数字冒充图表显示值。
    for node in tree.iter():
        tag = etree.QName(node).localname
        if tag in {"grouping", "barDir", "numFmt", "orientation", "min", "max"}:
            value = node.get("val", node.get("formatCode", ""))
            if value:
                lines.append(tag + "：" + value)
    series = charts[0].findall("{*}ser")
    if not series or len(series) > 64:
        raise HybridPreflightError("HYBRID_UNSUPPORTED_CHART")
    for index, item in enumerate(series, 1):
        lines.append(f"系列 {index}")
        for role in ("tx", "cat", "val", "xVal", "yVal"):
            element = item.find(f"{{*}}{role}")
            if element is not None and element.find(".//{*}f") is None:
                literal = [n.text or "" for n in element.findall(".//{*}v")]
                if len(literal) > 4096:
                    raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
                if literal:
                    lines.append(role + "：" + json.dumps(literal, ensure_ascii=False))
        refs = item.findall(".//{*}f")
        if not refs and not item.findall(".//{*}v"):
            raise HybridPreflightError("HYBRID_UNSUPPORTED_CHART")
        for ref in refs:
            expression = ref.text or ""
            match = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^'!]+))!([$A-Za-z0-9:]+)", expression)
            if not match:
                raise HybridPreflightError("HYBRID_UNSUPPORTED_CHART")
            sheet = (match[1] or match[2]).replace("''", "'")
            if sheet not in cell_values:
                raise HybridPreflightError("HYBRID_EXTERNAL_ASSET_MISSING")
            x0, y0, x1, y1 = range_boundaries(match[3])
            if None in (x0, y0, x1, y1) or (x1-x0+1)*(y1-y0+1) > 4096:
                raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
            if not (1 <= x0 <= x1 <= 16384 and 1 <= y0 <= y1 <= 1048576):
                raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
            # 返回引用及单元格原值；公式保存原式，不读取缓存结果、不执行表达式。
            from openpyxl.utils.cell import get_column_letter

            values = [cell_values[sheet].get(f"{get_column_letter(x)}{y}", "")
                      for y in range(y0, y1+1) for x in range(x0, x1+1)]
            lines.append("引用：" + expression + "\n值：" + json.dumps(values, ensure_ascii=False))
    return "\n".join(lines)


def sheet_cells(tree, strings: list[str]) -> dict[str, str]:
    """读取有界图表引用原值；共享公式用地址平移还原原式，不计算缓存结果。"""

    cells = tree.findall(".//{*}sheetData/{*}row/{*}c")
    if len(cells) > 1000000:
        raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
    result = {}
    shared_formulas = {}
    for cell in cells:
        address = cell.get("r", "")
        if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", address) or address in result:
            raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
        formula = cell.find("{*}f")
        value = cell.findtext("{*}v", "")
        if formula is not None:
            if formula.text:
                value = "=" + formula.text
                if formula.get("t") == "shared":
                    shared_formulas[formula.get("si")] = (address, value)
            elif formula.get("t") == "shared" and formula.get("si") in shared_formulas:
                origin, original = shared_formulas[formula.get("si")]
                value = Translator(original, origin=origin).translate_formula(address)
            else:
                raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
        elif cell.get("t") == "s":
            if not value.isdecimal() or int(value) >= len(strings):
                raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
            value = strings[int(value)]
        elif cell.get("t") == "inlineStr":
            value = "".join(node.text or "" for node in cell.findall(".//{*}t"))
        result[address] = value
    return result


def build_bundle(images: list[bytes], policy: HybridPolicy) -> bytes:
    """运行时把唯一知识图片整页嵌入一份 PDF，不绘制 OCR 标记，不逐图请求 API。"""

    if not images:
        return b""
    if len(images) > policy.max_pages:
        raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
    try:
        with pymupdf.open() as document:
            for content in images:
                width, height, _ = image_dimensions(content)
                scale = min(1.0, 1440 / max(width, height))
                page = document.new_page(width=max(3, width*scale), height=max(3, height*scale))
                page.insert_image(page.rect, stream=content)
            result = document.tobytes(deflate=True, no_new_id=True)
    except (ValueError, RuntimeError):
        raise HybridPreflightError() from None
    if len(result) > policy.max_bundle_bytes:
        raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
    return result


def prepare_hybrid(doc_id: UUID, path: Path, file_name: str, policy: HybridPolicy,
                   visual_lookup: Callable[[str], list[ParsedBlock] | None] | None = None) -> HybridPlan:
    """读取主结果和候选；过滤后先查精确视觉缓存，仅未命中图片进入一份 Bundle。"""

    try:
        native = XlsxParser().parse(doc_id=doc_id, file_path=path, file_name=file_name, file_type="xlsx")
    except NoTextError:
        # 视觉专用工作表可以没有单元格；后续必须确有视觉结果，不能把全空 Workbook 判成功。
        native = ParsedDocument(doc_id=doc_id, file_name=file_name, file_type="xlsx", blocks=[])
    entries, images, charts, cached = [], [], {}, {}
    seen_hashes, used_media, used_drawings, used_charts = {}, set(), set(), set()
    try:
        with ZipFile(path) as archive:
            names = archive.namelist()
            if any(n.startswith("xl/externalLinks/") for n in names):
                raise HybridPreflightError("HYBRID_EXTERNAL_ASSET_MISSING")
            if any(n.startswith(("xl/embeddings/", "xl/chartsheets/"))
                   or n.lower().endswith(".vml") for n in names):
                raise HybridPreflightError("HYBRID_UNSUPPORTED_SHEET")
            workbook = xml_part(archive, "xl/workbook.xml", policy.xml_limit)
            rels = relationships(archive, "xl/workbook.xml", policy.xml_limit)
            strings = []
            if "xl/sharedStrings.xml" in names:
                strings = ["".join(n.text or "" for n in item.findall(".//{*}t"))
                           for item in xml_part(archive, "xl/sharedStrings.xml", policy.xml_limit)]
            sheets, values = [], {}
            for sheet in workbook.findall(".//{*}sheet"):
                name = sheet.get("name", "")
                part = rels.get(sheet.get(f"{{{R}}}id"))
                if not name or name in values or not part or not part.startswith("xl/worksheets/"):
                    raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
                tree = xml_part(archive, part, policy.xml_limit)
                if any(etree.QName(n).localname in {"legacyDrawing", "legacyDrawingHF", "oleObjects", "controls", "picture", "extLst"}
                       for n in tree.iter()):
                    raise HybridPreflightError("HYBRID_UNSUPPORTED_SHEET")
                sheets.append((name, part, tree))
                values[name] = sheet_cells(tree, strings)
            if len(sheets) > 1000 or sum(len(v) for v in values.values()) > 1000000:
                raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
            for sheet_index, (name, part, tree) in enumerate(sheets, 1):
                sheet_rels = relationships(archive, part, policy.xml_limit)
                for drawing_ref in tree.findall("{*}drawing"):
                    drawing_part = sheet_rels.get(drawing_ref.get(f"{{{R}}}id"))
                    if not drawing_part or drawing_part in used_drawings:
                        raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
                    used_drawings.add(drawing_part)
                    drawing = xml_part(archive, drawing_part, policy.xml_limit)
                    drawing_rels = relationships(archive, drawing_part, policy.xml_limit)
                    for position, node in enumerate(drawing, 1):
                        if len(entries) >= policy.max_candidates:
                            raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
                        anchor = anchor_coordinates(node)
                        if any(etree.QName(n).localname not in {
                            "from", "to", "ext", "pic", "graphicFrame", "clientData"
                        } for n in node):
                            raise HybridPreflightError("HYBRID_UNSUPPORTED_DRAWING")
                        pictures = node.findall(f"{{{XDR}}}pic")
                        frames = node.findall(f"{{{XDR}}}graphicFrame")
                        if len(pictures) + len(frames) != 1 or any(
                            etree.QName(n).localname in {"sp", "grpSp", "cxnSp", "contentPart", "extLst"}
                            for n in node.iter()
                        ):
                            raise HybridPreflightError("HYBRID_UNSUPPORTED_DRAWING")
                        visual_id = uuid5(doc_id, f"{VERSION}:{sheet_index}:{drawing_part}:{position}:{anchor}")
                        if frames:
                            chart_refs = frames[0].findall(f".//{{{CHART}}}chart")
                            if len(chart_refs) != 1:
                                raise HybridPreflightError("HYBRID_UNSUPPORTED_DRAWING")
                            chart_part = drawing_rels.get(chart_refs[0].get(f"{{{R}}}id"))
                            if not chart_part:
                                raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
                            used_charts.add(chart_part)
                            text = chart_text(xml_part(archive, chart_part, policy.xml_limit), values)
                            # 标准 Chart Style/ColorStyle 是附属元数据，不把它们误当孤立 Chart。
                            for style_part in relationships(archive, chart_part, policy.xml_limit).values():
                                style_tree = xml_part(archive, style_part, policy.xml_limit)
                                if etree.QName(style_tree).localname not in {"chartStyle", "colorStyle"}:
                                    raise HybridPreflightError("HYBRID_UNSUPPORTED_CHART")
                                nodes = list(style_tree.iter())
                                if len(nodes) > 512:
                                    raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
                                text += "\n图表样式元数据：" + json.dumps([
                                    {"type": n.tag, "attributes": dict(sorted(n.attrib.items())),
                                     "text": n.text or ""} for n in nodes
                                ], ensure_ascii=False, sort_keys=True)
                                used_charts.add(style_part)
                            block = ParsedBlock(block_id=uuid5(doc_id, f"{VERSION}:{visual_id}:chart"),
                                                page=1, section=name, block_type="text", text=text,
                                                source="native_parser", bbox=None, confidence=1.0)
                            charts[visual_id] = block
                            entries.append(VisualEntry(visual_id=visual_id, source_doc_id=doc_id,
                                source_type="chart", source_section=name, sheet_index=sheet_index,
                                source_anchor=anchor, position=position, content_hash=digest(text.encode()),
                                media_type="application/chart", state="parsed", block_ids=[block.block_id]))
                            continue
                        picture = pictures[0]
                        blips = picture.findall(f".//{{{IMAGE}}}blip")
                        if len(blips) != 1:
                            raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
                        if blips[0].get(f"{{{R}}}link"):
                            raise HybridPreflightError("HYBRID_EXTERNAL_ASSET_MISSING")
                        if len(blips[0]):
                            raise HybridPreflightError("HYBRID_UNSUPPORTED_IMAGE_TRANSFORM")
                        # 空 effectLst 和全零 srcRect 是常见无操作结构，不应因存在标签拒绝。
                        if any((etree.QName(n).localname == "srcRect"
                                and any(n.get(side, "0") != "0" for side in ("l", "t", "r", "b")))
                               or (etree.QName(n).localname in {"effectLst", "effectDag"} and len(n))
                               or etree.QName(n).localname == "tile"
                               or n.get("rot", "0") != "0" or n.get("flipH", "0") not in {"0", "false"}
                               or n.get("flipV", "0") not in {"0", "false"} for n in picture.iter()):
                            raise HybridPreflightError("HYBRID_UNSUPPORTED_IMAGE_TRANSFORM")
                        media = drawing_rels.get(blips[0].get(f"{{{R}}}embed"))
                        if not media or not media.startswith("xl/media/"):
                            raise HybridPreflightError("HYBRID_INVALID_PACKAGE")
                        used_media.add(media)
                        if archive.getinfo(media).file_size > policy.max_image_bytes:
                            raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
                        content = archive.read(media)
                        width, height, media_type = image_dimensions(content)
                        if not width or not height or width*height > policy.max_pixels:
                            raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
                        image_hash = digest(content)
                        entry = VisualEntry(visual_id=visual_id, source_doc_id=doc_id,
                            source_type="embedded_image", source_section=name, sheet_index=sheet_index,
                            source_anchor=anchor, position=position, content_hash=image_hash,
                            media_type=media_type, width=width, height=height, state="pending")
                        # 只过滤可确定的微小/纯色图片；不靠 OCR 或文件名猜测业务语义。
                        pixmap = pymupdf.Pixmap(content)
                        if pixmap.width != width or pixmap.height != height:
                            raise HybridPreflightError("HYBRID_INVALID_IMAGE")
                        if min(width, height) < policy.min_side:
                            entry.state, entry.filter_reason = "filtered", "small"
                        elif is_solid(pixmap):
                            entry.state, entry.filter_reason = "filtered", "solid"
                        elif image_hash in seen_hashes:
                            previous = seen_hashes[image_hash]
                            entry.state, entry.reused_from = "reused", previous.visual_id
                            entry.bundle_page = previous.bundle_page
                            entry.cache_hit = previous.cache_hit
                            entry.cache_result_hash = previous.cache_result_hash
                        else:
                            reused = visual_lookup(image_hash) if visual_lookup is not None else None
                            if reused is not None:
                                cached[image_hash] = reused
                                entry.state, entry.cache_hit = "reused", True
                                entry.cache_result_hash = digest(json.dumps(
                                    [b.model_dump(mode="json") for b in reused], sort_keys=True,
                                    ensure_ascii=False).encode())
                                seen_hashes[image_hash] = entry
                                entries.append(entry)
                                continue
                            if len(images) >= policy.max_pages:
                                raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
                            if sum(len(image) for image in images) + len(content) > policy.max_bundle_bytes:
                                raise HybridPreflightError("HYBRID_RESOURCE_LIMIT")
                            images.append(content)
                            entry.bundle_page = len(images)
                            seen_hashes[image_hash] = entry
                        entries.append(entry)
            # 孤立媒体或未知绘图不能假定无意义并偷偷忽略。
            expected = ({n for n in names if n.startswith("xl/media/") and not n.endswith("/")},
                        {n for n in names if n.startswith("xl/drawings/") and n.endswith(".xml")},
                        {n for n in names if n.startswith("xl/charts/") and n.endswith(".xml")})
            if expected != (used_media, used_drawings, used_charts):
                raise HybridPreflightError("HYBRID_ORPHAN_PART")
            if not native.blocks and (
                not entries or any(value.strip() for cells in values.values() for value in cells.values())
            ):
                raise NoTextError("XLSX 缺少可用主结果或视觉内容")
            bundle = build_bundle(images, policy)
            manifest = BundleManifest(policy=policy, sheet_names=[s[0] for s in sheets], entries=entries,
                                      bundle_page_count=len(images), bundle_hash=digest(bundle),
                                      native_hash=digest(native.model_dump_json().encode("utf-8")))
            return HybridPlan(native, manifest, bundle, charts, cached)
    except (ParserError, OSError):
        raise
    except (BadZipFile, KeyError, ValueError, TypeError, AttributeError, RuntimeError, etree.XMLSyntaxError):
        raise HybridPreflightError("HYBRID_INVALID_PACKAGE") from None


def merge_hybrid(plan: HybridPlan, visual: ParsedDocument | None) -> ParsedDocument:
    """按Sheet/Anchor保留Native原式与已映射视觉块；质量继续传递，不掩盖云结果降级。"""

    manifest, native = plan.manifest, plan.native
    if (digest(native.model_dump_json().encode("utf-8")) != manifest.native_hash
        or len(set(manifest.sheet_names)) != len(manifest.sheet_names)
        or any(e.source_section not in manifest.sheet_names for e in manifest.entries)):
        raise HybridPreflightError()
    if len({entry.visual_id for entry in manifest.entries}) != len(manifest.entries):
        raise HybridPreflightError()
    pages = {}
    missing_visuals = set()
    if manifest.bundle_page_count:
        if (visual is None or visual.doc_id != native.doc_id or visual.file_type != "pdf"
            or visual.file_name != native.file_name):
            raise HybridPreflightError()
        for block in visual.blocks:
            if block.source != "mineru_cloud" or not block.text.strip():
                raise HybridPreflightError()
            pages.setdefault(block.page, []).append(block)
        expected_pages = set(range(1, manifest.bundle_page_count+1))
        partial_visual = visual.quality is not None and visual.quality.content == "PARTIAL"
        if not set(pages) <= expected_pages or set(pages) != expected_pages and not partial_visual:
            raise HybridPreflightError()
    elif visual is not None:
        raise HybridPreflightError()
    if any(block.source != "native_parser" or block.page != 1 or block.bbox is not None
           or block.block_type != "table" for block in native.blocks):
        raise HybridPreflightError()
    if len({block.section for block in native.blocks}) != len(native.blocks):
        raise HybridPreflightError()
    blocks = []
    for sheet in manifest.sheet_names:
        # Native 粒度是整 Sheet；不拆表伪造逐单元格插入，增强项放在同 Sheet table 后。
        blocks.extend(block for block in native.blocks if block.section == sheet)
        ordered = sorted((e for e in manifest.entries if e.source_section == sheet),
                         key=lambda e: (e.source_anchor[0], e.source_anchor[1], e.position))
        for entry in ordered:
            if (entry.source_doc_id != native.doc_id or len(entry.source_anchor) not in {4, 8}
                or entry.sheet_index != manifest.sheet_names.index(sheet) + 1):
                raise HybridPreflightError()
            if entry.state == "filtered":
                if entry.filter_reason not in {"small", "solid"} or entry.bundle_page is not None:
                    raise HybridPreflightError()
                continue
            if entry.source_type == "chart":
                block = plan.charts.get(entry.visual_id)
                if block is None or digest(block.text.encode()) != entry.content_hash:
                    raise HybridPreflightError()
                blocks.append(block)
                continue
            if entry.state not in {"pending", "parsed", "reused"}:
                raise HybridPreflightError()
            if entry.cache_hit:
                source_blocks = plan.cached.get(entry.content_hash)
                if (entry.bundle_page is not None or not source_blocks
                    or any(b.source != "mineru_cloud" or not b.text.strip() or b.page != 1
                           or b.bbox is not None or b.section for b in source_blocks)):
                    raise HybridPreflightError()
            else:
                if entry.bundle_page not in pages:
                    # 原Bundle来源已由manifest核对，未识别视觉页不补造正文；Native原式仍保留。
                    if visual is None or visual.quality is None or visual.quality.content != "PARTIAL":
                        raise HybridPreflightError()
                    entry.state, entry.block_ids = "failed", []
                    missing_visuals.add(entry.visual_id)
                    continue
                source_blocks = pages[entry.bundle_page]
            if entry.reused_from is not None:
                original = next((e for e in manifest.entries if e.visual_id == entry.reused_from), None)
                if (original is None or original.reused_from is not None
                    or original.content_hash != entry.content_hash or original.bundle_page != entry.bundle_page
                    or original.cache_hit != entry.cache_hit):
                    raise HybridPreflightError()
            entry.block_ids = []
            seen = set()
            for index, block in enumerate(source_blocks, 1):
                key = (block.block_type, " ".join(block.text.split()))
                if key in seen:
                    continue
                seen.add(key)
                copied = block.model_copy(update={
                    "block_id": uuid5(native.doc_id, f"{VERSION}:{entry.visual_id}:{index}"),
                    "section": sheet, "page": 1, "bbox": None,
                    "provenance": block.provenance.model_copy(update={"bbox_reliable": False,
                        "pointer": f"hybrid/{entry.visual_id}/" + block.provenance.pointer}) if block.provenance else None,
                })
                blocks.append(copied)
                entry.block_ids.append(copied.block_id)
            entry.state = "reused" if entry.reused_from is not None or entry.cache_hit else "parsed"
            if not entry.block_ids:
                raise HybridPreflightError()
    if any(e.state not in {"parsed", "filtered", "reused"} and e.visual_id not in missing_visuals for e in manifest.entries):
        raise HybridPreflightError()
    if any(block.section not in manifest.sheet_names for block in native.blocks):
        raise HybridPreflightError()
    if not blocks or not any(block.text.strip() for block in blocks):
        raise NoTextError("XLSX 过滤后没有有效正文，不能判为解析成功")
    quality = visual.quality.model_copy(deep=True) if visual and visual.quality else None
    if quality:
        # 对外沿用原XLSX逻辑页，不能把Bundle中间页当成原件页码。
        quality.covered_pages, quality.required_pages, quality.missing_pages = [1], [1], []
        if missing_visuals and "MISSING_VISUAL_PAGES" not in quality.warnings and len(quality.warnings) < 128:
            quality.warnings.append("MISSING_VISUAL_PAGES")
    return ParsedDocument(doc_id=native.doc_id, file_name=native.file_name, file_type="xlsx",
                          title=native.title, blocks=blocks, quality=quality,
                          quarantined_content=visual.quarantined_content if visual else [])


def is_solid(pixmap) -> bool:
    """只认定所有像素完全相同的纯色图；黑白扫描文字不能当成纯色装饰图。"""

    pixels = pixmap.samples
    return bool(pixels) and pixels == pixels[:pixmap.n] * (pixmap.width * pixmap.height)

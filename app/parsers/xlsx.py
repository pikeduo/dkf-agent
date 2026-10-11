"""只提取 XLSX 单元格，保留工作表顺序与公式文本，不解析视觉对象或执行公式。"""

import math
import re
from datetime import date, datetime, time
from pathlib import Path
from uuid import UUID
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from lxml.etree import XMLSyntaxError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.file_formats import FileFormatError, validate_file_content
from app.parsers.base import BaseParser
from app.parsers.common import NoTextError, ParserError, make_block
from app.parsers.csv import table_text
from app.schemas.parsed_document import FileType, ParsedDocument


def cell_text(cell) -> str:
    """保留公式及标量值；普通百分比格式可读化，其余自定义格式不模拟 Excel 渲染。"""

    value = cell.value
    if value is None:
        return ""
    if cell.data_type == "f":
        formula = value if isinstance(value, str) else getattr(value, "text", None)
        if not isinstance(formula, str) or not formula:
            raise ParserError("XLSX 包含无法读取文本的特殊公式")
        return formula
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise ParserError("XLSX 包含非有限数值")
        percent = re.fullmatch(r"0(?:\.(0{1,12}))?%", cell.number_format or "")
        if percent:
            return f"{value * 100:.{len(percent[1] or '')}f}%"
        return str(value)
    if isinstance(value, str):
        return value
    raise ParserError("XLSX 包含尚未支持的单元格结构")


class XlsxParser(BaseParser):
    """每个非空工作表形成一个 table 块，section 为 Sheet 名，page 为逻辑页 1。"""

    def _parse(
        self, *, doc_id: UUID, file_path: Path, file_name: str, file_type: FileType
    ) -> ParsedDocument:
        """只读打开 OOXML，逐行读取原值与公式；关闭资源后返回稳定块，不触发云端路由。"""

        if file_type != "xlsx":
            raise ParserError("XLSX 解析器仅接受 xlsx 文件，旧 XLS 请显式提交 MinerU")
        workbook = None
        try:
            with file_path.open("rb") as stream:
                validate_file_content(stream, file_type)
            # data_only=False 保留公式而不是上次 Excel 缓存值；不加载外部工作簿链接。
            workbook = load_workbook(
                file_path, read_only=True, data_only=False, keep_links=False
            )
            if len(workbook.worksheets) > 1000:
                raise ParserError("XLSX 工作表数量超过安全上限")
            blocks = []
            cells = 0
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                if (sheet.max_row or 0) > 100000 or (sheet.max_column or 0) > 1024:
                    raise ParserError("XLSX 行列规模超过原生解析安全上限")
                # 不依赖可能低估真实范围的 dimension 标签，按 XML 实际行列读取并再次限量。
                sheet.reset_dimensions()
                rows = []
                width = 0
                last_nonempty = 0
                for index, row in enumerate(sheet.iter_rows(), start=1):
                    cells += len(row)
                    if index > 100000 or len(row) > 1024 or cells > 1000000:
                        raise ParserError("XLSX 单元格规模超过原生解析安全上限")
                    values = [cell_text(cell) for cell in row]
                    nonempty = [
                        column
                        for column, value in enumerate(values, start=1)
                        if value.strip()
                    ]
                    if nonempty:
                        width = max(width, nonempty[-1])
                        last_nonempty = index
                    rows.append(values)
                if not last_nonempty:
                    continue
                ordered_rows = [
                    row[:width] + [""] * max(0, width - len(row))
                    for row in rows[:last_nonempty]
                ]
                blocks.append(
                    make_block(
                        doc_id,
                        sheet_index,
                        page=1,
                        block_type="table",
                        text=table_text(ordered_rows),
                        section=sheet.title,
                    )
                )
            if not blocks:
                raise NoTextError("XLSX 所有工作表均无可用单元格内容")
            return ParsedDocument(
                doc_id=doc_id, file_name=file_name, file_type=file_type, blocks=blocks
            )
        except FileFormatError:
            raise ParserError("XLSX 容器格式损坏或与扩展名不匹配") from None
        except (ParserError, OSError):
            raise
        except (
            BadZipFile,
            InvalidFileException,
            ParseError,
            XMLSyntaxError,
            ValueError,
            TypeError,
            KeyError,
            IndexError,
            AttributeError,
            RuntimeError,
        ):
            # 第三方读取错误可能含路径或原件内容，不将其原文暴露给任务结果。
            raise ParserError(
                "XLSX 工作簿无法读取，请检查损坏、加密或不支持的结构"
            ) from None
        finally:
            if workbook is not None:
                workbook.close()

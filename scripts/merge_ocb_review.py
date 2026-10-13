#!/usr/bin/env python3
"""Merge manually approved OCB Office source reviews without fabricating permissions.

Place alongside the existing DKF-Agent project files:
  scripts/merge_ocb_review.py
  docs/test/ocb_preset/results/ocb_license_review.xlsx (or .csv)
  docs/test/ocb_preset/results/license_review_candidates.json
  docs/test/ocb_preset/results/native_download_plan.json
  docs/test/ocb_preset/source/source_review.json

Default is PREVIEW ONLY. --apply is required to write source_review.json.
No networking, downloading, parsing, or cloud calls.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import sys
from pathlib import Path
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/test/ocb_preset/source"
RESULTS = ROOT / "docs/test/ocb_preset/results"
FIELDS = (
    "reference_file", "file_type", "source_kind", "source_url",
    "review_decision", "license", "license_evidence", "original_native",
    "local_use_confirmed", "allowed_redirect_hosts", "original_format",
    "required_format", "conversion_required", "qa_count", "candidate_id",
)
OUTPUT_FIELDS = (
    "allowed_redirect_hosts", "conversion_required", "license",
    "local_use_confirmed", "original_native", "reference_file", "source_url",
)


class ReviewError(Exception):
    pass


def load_json(path: Path):
    try:
        with path.open("r", encoding="utf-8-sig") as stream:
            return json.load(stream)
    except (OSError, ValueError) as exc:
        raise ReviewError(f"无法读取合法 JSON：{path} ({type(exc).__name__})") from None


def check_plan(candidates: dict, plan_path: Path):
    if not isinstance(candidates, dict) or not isinstance(candidates.get("candidates"), list):
        raise ReviewError("license_review_candidates.json 的顶层结构不正确")
    revision = candidates.get("source", {}).get("revision")
    if not isinstance(revision, str) or not re.fullmatch(r"[0-9a-f]{40}", revision):
        raise ReviewError("候选来源 revision 无效")
    if not plan_path.is_file():
        raise ReviewError(f"未找到下载计划：{plan_path}")
    plan = load_json(plan_path)
    if candidates.get("plan_sha256") != plan.get("plan_sha256"):
        raise ReviewError("两个 JSON 计划摘要不一致；请重新生成匹配的审核表")
    keys = set()
    files = set()
    result = {}
    by_id = {obj.get("candidate_id"): obj for obj in plan.get("candidates", [])}
    for row in candidates["candidates"]:
        if not isinstance(row, dict) or not isinstance(row.get("review_template"), dict):
            raise ReviewError("存在非对象或缺少 review_template 的候选")
        ident, name = row.get("candidate_id"), row.get("reference_file")
        template = row["review_template"]
        if (not isinstance(ident, str) or not isinstance(name, str)
                or ident in keys or name in files):
            raise ReviewError("候选 ID 或 reference_file 无效/重复")
        if (template.get("reference_file") != name
                or template.get("source_url") != row.get("source_url")
                or template.get("conversion_required") is not False):
            raise ReviewError("候选模板的原始来源信息不一致")
        original = by_id.get(ident)
        if (not original or original.get("reference_file") != name
                or original.get("source_url") != row.get("source_url")
                or original.get("file_type") != row.get("file_type")
                or original.get("conversion_required") is not False):
            raise ReviewError("审核候选与下载计划不一致")
        keys.add(ident)
        files.add(name)
        result[ident] = row
    return result


def get_rows(sheet: Path):
    if sheet.suffix.lower() == ".csv":
        try:
            with sheet.open("r", encoding="utf-8-sig", newline="") as stream:
                reader = csv.DictReader(stream)
                if not reader.fieldnames or set(FIELDS) - set(reader.fieldnames):
                    raise ReviewError("CSV 表头缺少必要字段，请使用配套的审核表")
                return list(reader)
        except (OSError, UnicodeError, csv.Error) as exc:
            raise ReviewError(f"CSV 读取失败：{type(exc).__name__}") from None
    if sheet.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ReviewError("未安装 openpyxl。请改用配套 CSV 审核表并指定 --sheet 路径") from None
        try:
            book = load_workbook(sheet, read_only=True, data_only=True)
            if "审核清单" not in book.sheetnames:
                raise ReviewError("Excel 缺少‘审核清单’工作表")
            rows = book["审核清单"].iter_rows(values_only=True)
            headings = next(rows, ())
            if set(FIELDS) - set(headings):
                raise ReviewError("Excel 表头缺少必要字段，请使用配套审核表")
            return [dict(zip(headings, values)) for values in rows if any(x is not None for x in values)]
        except (OSError, ValueError, KeyError) as exc:
            raise ReviewError(f"Excel 读取失败：{type(exc).__name__}") from None
    raise ReviewError("--sheet 只支持 .xlsx 或 .csv")


def text(value) -> str:
    return "" if value is None else str(value).strip()


def truth(value) -> bool:
    return value is True or (isinstance(value, str) and value.strip().casefold() == "true")


def falsehood(value) -> bool:
    return value is False or (isinstance(value, str) and value.strip().casefold() == "false")


def redirect_hosts(value):
    if value is None or text(value) == "":
        return []
    try:
        hosts = json.loads(value) if isinstance(value, str) else value
    except (TypeError, ValueError):
        raise ReviewError("allowed_redirect_hosts 必须为 JSON 数组，例如 [] 或 [\"cdn.example.org\"]") from None
    if not isinstance(hosts, list) or len(hosts) > 16:
        raise ReviewError("allowed_redirect_hosts 必须为长度不超过 16 的数组")
    for host in hosts:
        if (not isinstance(host, str) or not host or len(host) > 253
                or not re.fullmatch(r"[A-Za-z0-9.-]+", host)
                or ".." in host or host.startswith(".") or host.endswith(".")):
            raise ReviewError("allowed_redirect_hosts 仅允许经过人工核对的纯主机名")
    if len({host.casefold() for host in hosts}) != len(hosts):
        raise ReviewError("allowed_redirect_hosts 有重复域名")
    return hosts


def validated_approvals(rows, by_id):
    approved, pending, rejected, seen_ids, seen_names = [], 0, 0, set(), set()
    for row in rows:
        ident = text(row.get("candidate_id"))
        reference = by_id.get(ident)
        if reference is None or ident in seen_ids:
            raise ReviewError(f"审核表有未知或重复的 candidate_id：{ident[:24]}")
        seen_ids.add(ident)
        name = reference["reference_file"]
        if name in seen_names:
            raise ReviewError(f"审核表存在重复文件：{name}")
        seen_names.add(name)
        template = reference["review_template"]
        fixed = {
            "reference_file": name,
            "file_type": reference["file_type"],
            "source_kind": reference["source_kind"],
            "source_url": reference["source_url"],
            "original_format": reference["native_criteria"]["original_format"],
            "required_format": reference["native_criteria"]["required_format"],
        }
        for key, expected in fixed.items():
            if text(row.get(key)) != text(expected):
                raise ReviewError(f"第 {len(seen_ids) + 1} 行 {key} 被修改：{name}；请使用未修改的候选元数据")
        if not falsehood(row.get("conversion_required")):
            raise ReviewError(f"第 {len(seen_ids) + 1} 行 conversion_required 必须保持 FALSE")
        decision = text(row.get("review_decision")).casefold()
        if decision in ("", "pending", "待审"):
            pending += 1
            continue
        if decision in ("rejected", "不通过"):
            rejected += 1
            continue
        if decision not in ("approved", "通过"):
            raise ReviewError(f"审核状态错误：{name}（使用 pending/approved/rejected）")
        license_name = text(row.get("license"))
        evidence = text(row.get("license_evidence"))
        if (not license_name or license_name.casefold() in ("unknown", "pending", "tbd", "none", "null", "待审", "未确认", "未知")
                or not evidence):
            raise ReviewError(f"已标记 approved 但许可证或许可依据为空/未知：{name}")
        if not truth(row.get("original_native")) or not truth(row.get("local_use_confirmed")):
            raise ReviewError(f"已标记 approved 但原生格式或本地使用未经确认：{name}")
        url = text(row.get("source_url"))
        if urlsplit(url).scheme != "https" or not urlsplit(url).hostname:
            raise ReviewError(f"已批准但 source_url 非合法 HTTPS URL：{name}")
        item = {
            "reference_file": name,
            "source_url": url,
            "license": license_name,
            "original_native": True,
            "conversion_required": False,
            "local_use_confirmed": True,
            "allowed_redirect_hosts": redirect_hosts(row.get("allowed_redirect_hosts")),
        }
        if item["source_url"] != template["source_url"]:
            raise ReviewError(f"审核表 source_url 不匹配：{name}")
        approved.append(item)
    if len(seen_ids) != len(by_id):
        print(f"提醒：审核表有 {len(by_id) - len(seen_ids)} 条原候选未列出，保持未审核。")
    return approved, pending, rejected


def validated_existing(path: Path):
    if not path.is_file():
        return [], b""
    raw = path.read_bytes()
    existing = load_json(path)
    if not isinstance(existing, list):
        raise ReviewError("已有 source_review.json 不是 JSON 数组，拒绝覆盖")
    seen = set()
    for obj in existing:
        name = obj.get("reference_file") if isinstance(obj, dict) else None
        if not isinstance(name, str) or not name or name in seen:
            raise ReviewError("已有 source_review.json 结构非法或有重复记录，拒绝覆盖")
        seen.add(name)
    return existing, raw


def write_safe(path: Path, raw: bytes, backup_dir: Path, new: list):
    if not path.parent.is_dir() or not backup_dir.is_dir():
        raise ReviewError("目标 source/results 目录不存在，停止写入")
    if path.is_symlink():
        raise ReviewError("目标是链接，禁止写入")
    if path.is_file() and path.read_bytes() != raw:
        raise ReviewError("审核期间 source_review.json 已变化，请重新预览")
    if raw:
        timestamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        backup = backup_dir / f"source_review_backup_{timestamp}.json"
        with backup.open("xb") as stream:
            stream.write(raw)
            stream.flush()
            os.fsync(stream.fileno())
        print(f"已备份原审核记录：{backup}")
    content = (json.dumps(new, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8")
    temp = path.with_name(path.name + f".tmp-{os.getpid()}")
    with temp.open("xb") as stream:
        stream.write(content)
        stream.flush()
        os.fsync(stream.fileno())
    try:
        os.replace(temp, path)
    finally:
        if temp.exists():
            temp.unlink()
    print(f"已合并：{path}")


def main(argv=None):
    parser = argparse.ArgumentParser(description="审核后的 OCB Office 许可记录安全合并（默认只预览）")
    parser.add_argument("--sheet", type=Path, default=RESULTS / "ocb_license_review.xlsx",
                        help="配套的 Excel/CSV 审核表")
    parser.add_argument("--candidates", type=Path, default=RESULTS / "license_review_candidates.json")
    parser.add_argument("--plan", type=Path, default=RESULTS / "native_download_plan.json")
    parser.add_argument("--target", type=Path, default=SOURCE / "source_review.json")
    parser.add_argument("--apply", action="store_true", help="明确确认后备份并合并到 source_review.json")
    args = parser.parse_args(argv)
    try:
        data = load_json(args.candidates)
        candidate_index = check_plan(data, args.plan)
        rows = get_rows(args.sheet)
        approved, pending, rejected = validated_approvals(rows, candidate_index)
        existing, original_bytes = validated_existing(args.target)
        indexed = {x["reference_file"]: x for x in existing}
        additions, same = [], 0
        for row in approved:
            old = indexed.get(row["reference_file"])
            if old is None:
                additions.append(row)
            elif all(old.get(key, [] if key == "allowed_redirect_hosts" else None) == value
                     for key, value in row.items()):
                same += 1
            else:
                raise ReviewError(f"现有审核与新审核冲突（不自动覆盖）：{row['reference_file']}")
        print(f"候选总数={len(candidate_index)}，表格行数={len(rows)}，已批准={len(approved)}，待审={pending}，拒绝={rejected}")
        print(f"现有审核={len(existing)}，准备新增={len(additions)}，相同跳过={same}")
        if not args.apply:
            print("当前为只读预览。检查无误后加 --apply 才写入 source_review.json。")
            return 0
        if not additions:
            print("没有新增审核记录，未修改 source_review.json。")
            return 0
        write_safe(args.target, original_bytes, args.target.parent.parent / "results", existing + additions)
        print("审核记录已写入。下载许可和云端外传许可仍需分别遵守实际条款。")
        return 0
    except ReviewError as exc:
        print("REVIEW_ERROR：" + str(exc), file=sys.stderr)
        return 2
    except OSError as exc:
        print("REVIEW_IO：" + type(exc).__name__, file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

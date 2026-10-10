"""原生解析与落库验收：仅使用隔离 schema、临时原件和本进程 Celery，不访问 Redis。"""

from io import BytesIO
from uuid import UUID, uuid4
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document as WordDocument
from sqlalchemy import func, inspect, select, text
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import Session

from alembic import command
from app.models import DocumentBlock, DocumentChunk
from app.services import document_processing as processing
from app.tasks import documents as tasks

pytest_plugins = ("document_processing_smoke",)
PREFIX = "/api/admin/knowledge-bases"


def image_content():
    """在内存生成覆盖页面的大幅 PNG，避免依赖外部样本或写入真实文件。"""

    pixmap = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 300, 300), 0)
    pixmap.clear_with(200)
    return pixmap.tobytes("png")


def pdf_content(kind="native"):
    """创建双页文本、扫描或混合 PDF；页码和图片均来自真实 PDF 结构。"""

    with pymupdf.open() as pdf:
        first = pdf.new_page(width=400, height=500)
        if kind in {"native", "mixed"}:
            first.insert_text((50, 60), "First page evidence.")
        if kind == "scan":
            first.insert_image(first.rect, stream=image_content())
        if kind in {"native", "mixed"}:
            second = pdf.new_page(width=400, height=500)
            if kind == "native":
                second.insert_text((50, 80), "Second page evidence.")
            else:
                second.insert_image(second.rect, stream=image_content())
        return pdf.tobytes()


def native_sample(file_type):
    """生成四种真实原生样本及预期阅读顺序，不将上传签名样本当作完整文档。"""

    if file_type == "txt":
        return "第一段正文。\n\n第二段正文。".encode(), ["第一段正文", "第二段正文"]
    if file_type == "md":
        content = "# 总标题\n\n第一段正文。\n\n## 二级章节\n\n第二段正文。"
        return content.encode(), ["总标题", "第一段正文", "二级章节", "第二段正文"]
    if file_type == "docx":
        word = WordDocument()
        word.add_heading("总标题", level=1)
        word.add_paragraph("第一段正文。")
        word.add_heading("二级章节", level=2)
        word.add_paragraph("第二段正文。")
        buffer = BytesIO()
        word.save(buffer)
        return buffer.getvalue(), ["总标题", "第一段正文", "二级章节", "第二段正文"]
    return pdf_content(), ["First page evidence", "Second page evidence"]


def upload_sample(context, file_name, content):
    """调用真实上传接口并由共享夹具模拟 Broker 接收，返回当前任务代次。"""

    api, kb_id, _ = context
    response = api(
        "POST",
        f"{PREFIX}/{kb_id}/documents",
        files={"file": (file_name, content, "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    return response.json()


def execute_document(document):
    """同步执行指定 Celery 消息，不发送消息、不保存真实 Redis 任务结果。"""

    return tasks.process_document.apply(
        args=[document["doc_id"]], task_id=document["task_id"], throw=False
    )


def document_state(context, document):
    """从真实详情接口读取任务提交后的数据库状态，不使用上传响应快照。"""

    api, kb_id, _ = context
    response = api("GET", f"{PREFIX}/{kb_id}/documents/{document['doc_id']}")
    assert response.status_code == 200, response.text
    return response.json()


def block_snapshot(migrated_connection, document):
    """按明确的原文序号读取所有块字段，以比较顺序、内容和重试稳定性。"""

    connection, _, _ = migrated_connection
    with Session(connection, join_transaction_mode="create_savepoint") as session:
        rows = session.scalars(
            select(DocumentBlock)
            .where(DocumentBlock.doc_id == UUID(document["doc_id"]))
            .order_by(DocumentBlock.block_index, DocumentBlock.block_id)
        )
        return [
            {
                "block_id": row.block_id,
                "doc_id": row.doc_id,
                "block_index": row.block_index,
                "page": row.page,
                "section": row.section,
                "block_type": row.block_type,
                "text": row.text,
                "bbox": row.bbox,
                "source": row.source,
                "confidence": row.confidence,
            }
            for row in rows
        ]


def assert_no_chunks(migrated_connection, document):
    """确认本阶段没有产生 Chunk，不将原生解析完成冒充完整知识库入库。"""

    connection, _, _ = migrated_connection
    with Session(connection, join_transaction_mode="create_savepoint") as session:
        assert (
            session.scalar(
                select(func.count())
                .select_from(DocumentChunk)
                .where(DocumentChunk.doc_id == UUID(document["doc_id"]))
            )
            == 0
        )


@pytest.mark.parametrize("file_type", ["txt", "md", "docx", "pdf"])
def test_native_upload_task_persists_ordered_uniform_blocks(
    processing_context, migrated_connection, monkeypatch, file_type
):
    """四种原生文档经真实上传和任务生成统一块，保留阅读顺序、标题和物理页码。"""

    content, expected = native_sample(file_type)
    document = upload_sample(processing_context, f"原生参考.{file_type}", content)
    result = execute_document(document)
    assert result.state == "SUCCESS", str(result.result)
    blocks = block_snapshot(migrated_connection, document)
    assert blocks
    assert result.result == {
        "doc_id": document["doc_id"],
        "document_status": "CHUNKING",
        "status": "awaiting_chunk",
        "block_count": len(blocks),
    }
    assert [block["block_index"] for block in blocks] == list(range(1, len(blocks) + 1))
    assert len({block["block_id"] for block in blocks}) == len(blocks)
    assert all(block["doc_id"] == UUID(document["doc_id"]) for block in blocks)
    assert all(block["source"] == "native_parser" for block in blocks)
    assert all(block["confidence"] == 1.0 for block in blocks)
    joined = "\n".join(block["text"] for block in blocks)
    positions = [joined.index(fragment) for fragment in expected]
    assert positions == sorted(positions)
    if file_type in {"md", "docx"}:
        assert [
            block["text"] for block in blocks if block["block_type"] == "title"
        ] == [
            "总标题",
            "二级章节",
        ]
    if file_type == "pdf":
        assert [block["page"] for block in blocks] == [1, 2]
        assert all(
            isinstance(block["bbox"], list) and len(block["bbox"]) == 4
            for block in blocks
        )
    else:
        assert all(block["page"] == 1 and block["bbox"] is None for block in blocks)
    detail = document_state(processing_context, document)
    assert detail["status"] == "CHUNKING" and detail["error_message"] is None
    assert detail["task_id"] == document["task_id"]

    def unexpected_reparse(file_type):
        """重复消息不得再次取得 Parser，否则说明幂等状态防护失效。"""

        raise AssertionError("重复任务触发了第二次解析")

    monkeypatch.setattr(processing, "get_parser", unexpected_reparse)
    duplicate = execute_document(document)
    assert duplicate.state == "SUCCESS"
    assert duplicate.result["status"] == "skipped"
    assert block_snapshot(migrated_connection, document) == blocks
    assert document_state(processing_context, document) == detail
    assert_no_chunks(migrated_connection, document)


@pytest.mark.parametrize("existing_blocks", [False, True])
def test_failed_block_write_rolls_back_everything_and_can_resubmit(
    processing_context, migrated_connection, monkeypatch, existing_blocks
):
    """块写入失败整体回滚并保留旧块，人工重投使用新代次且不会生成重复数据。"""

    content, _ = native_sample("txt")
    document = upload_sample(processing_context, "事务恢复.txt", content)
    connection, _, _ = migrated_connection
    if existing_blocks:
        with Session(connection, join_transaction_mode="create_savepoint") as session:
            session.add(
                DocumentBlock(
                    doc_id=UUID(document["doc_id"]),
                    block_index=1,
                    page=1,
                    block_type="text",
                    text="此前保存的解析块",
                    source="native_parser",
                    confidence=1.0,
                )
            )
            session.commit()
    before = block_snapshot(migrated_connection, document)
    original = processing.persist_blocks

    def fail_after_flush(session, row, parsed):
        """先实际删除旧块并刷新新块再抛错，验证回滚而不是仅测试解析前失败。"""

        original(session, row, parsed)
        session.flush()
        raise ValueError("private SQL secret_password")

    with monkeypatch.context() as patch:
        patch.setattr(processing, "persist_blocks", fail_after_flush)
        result = execute_document(document)
    assert result.state == "FAILURE"
    failed = document_state(processing_context, document)
    assert failed["status"] == "FAILED"
    assert "secret_password" not in failed["error_message"]
    assert block_snapshot(migrated_connection, document) == before
    assert_no_chunks(migrated_connection, document)
    api, kb_id, _ = processing_context
    response = api("POST", f"{PREFIX}/{kb_id}/documents/{document['doc_id']}/process")
    assert response.status_code == 202, response.text
    current = response.json()
    assert current["task_id"] != document["task_id"]
    assert current["error_message"] is None
    assert execute_document(document).result["status"] == "stale_task"
    assert execute_document(current).state == "SUCCESS"
    blocks = block_snapshot(migrated_connection, current)
    assert len(blocks) == 2
    assert [block["block_index"] for block in blocks] == [1, 2]
    assert not {block["block_id"] for block in before} & {
        block["block_id"] for block in blocks
    }
    snapshot = document_state(processing_context, current)
    tasks.save_failure(
        UUID(document["doc_id"]), UUID(document["task_id"]), "旧任务错误", True
    )
    assert document_state(processing_context, current) == snapshot
    assert block_snapshot(migrated_connection, current) == blocks


def test_transient_block_write_retries_without_partial_or_duplicate_rows(
    processing_context, migrated_connection, monkeypatch
):
    """数据库临时故障在真实写入后触发重试，回滚后恢复且块标识稳定、无重复。"""

    content, _ = native_sample("txt")
    document = upload_sample(processing_context, "数据库重试.txt", content)
    original = processing.persist_blocks
    attempts = []

    def transient_write(session, row, parsed):
        """第一次写入后模拟断连，下一次正常完成且使用相同解析块标识。"""

        original(session, row, parsed)
        attempts.append([block.block_id for block in parsed.blocks])
        if len(attempts) == 1:
            raise OperationalError("private SQL", None, Exception("secret_password"))

    monkeypatch.setattr(processing, "persist_blocks", transient_write)
    result = execute_document(document)
    assert result.state == "SUCCESS", str(result.result)
    assert len(attempts) == 2 and attempts[0] == attempts[1]
    blocks = block_snapshot(migrated_connection, document)
    assert len(blocks) == 2
    assert [block["block_id"] for block in blocks] == attempts[1]
    assert document_state(processing_context, document)["error_message"] is None
    assert_no_chunks(migrated_connection, document)


@pytest.mark.parametrize("kind", ["scan", "mixed"])
def test_pdf_requiring_ocr_fails_without_partial_native_blocks(
    processing_context, migrated_connection, kind
):
    """扫描和混合 PDF 不将部分原生文本冒充完整解析，也不启动未实现的 OCR。"""

    document = upload_sample(
        processing_context, f"需要识别-{kind}.pdf", pdf_content(kind)
    )
    result = execute_document(document)
    assert result.state == "FAILURE"
    detail = document_state(processing_context, document)
    assert detail["status"] == "FAILED"
    assert "OCR" in detail["error_message"]
    assert str(processing_context[2]) not in detail["error_message"]
    assert block_snapshot(migrated_connection, document) == []
    assert_no_chunks(migrated_connection, document)


def test_image_requires_ocr_and_is_not_fake_success(
    processing_context, migrated_connection
):
    """图片仍可上传，但当前原生任务清晰报告需要 OCR，不生成虚假的正文或成功状态。"""

    document = upload_sample(processing_context, "待识别.png", image_content())
    result = execute_document(document)
    assert result.state == "FAILURE"
    detail = document_state(processing_context, document)
    assert detail["status"] == "FAILED"
    assert "OCR" in detail["error_message"]
    assert block_snapshot(migrated_connection, document) == []
    assert_no_chunks(migrated_connection, document)


def invalid_content(file_type):
    """生成可通过上传签名校验但不能完整解析的样本，验证 Parser 的安全失败路径。"""

    if file_type == "txt":
        return b"\xff\xd0\x80"
    if file_type == "pdf":
        return b"%PDF-1.4\ninvalid document fixture\n%%EOF"
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
    return buffer.getvalue()


@pytest.mark.parametrize("file_type", ["txt", "docx", "pdf"])
def test_invalid_native_document_records_safe_failure_without_blocks(
    processing_context, migrated_connection, file_type
):
    """编码或文档结构错误永久失败，任务与数据库不泄漏文件路径且不保留半成品。"""

    document = upload_sample(
        processing_context, f"无效内容.{file_type}", invalid_content(file_type)
    )
    result = execute_document(document)
    assert result.state == "FAILURE"
    detail = document_state(processing_context, document)
    assert detail["status"] == "FAILED" and detail["error_message"]
    assert detail["error_message"] == str(result.result)
    assert str(processing_context[2]) not in str(result.result)
    assert block_snapshot(migrated_connection, document) == []
    assert_no_chunks(migrated_connection, document)


@pytest.mark.parametrize("block_index", [0, -1])
def test_block_index_rejects_nonpositive_values(
    processing_context, migrated_connection, block_index
):
    """数据库拒绝非正数原文序号，约束独立于 Parser 或 ORM 的输入检查。"""

    document = upload_sample(processing_context, "顺序约束.txt", b"block index fixture")
    connection, _, _ = migrated_connection
    with Session(connection, join_transaction_mode="create_savepoint") as session:
        with pytest.raises(IntegrityError), session.begin_nested():
            session.add(
                DocumentBlock(
                    doc_id=UUID(document["doc_id"]),
                    block_index=block_index,
                    page=1,
                    block_type="text",
                    text="非法顺序",
                    source="native_parser",
                )
            )
            session.flush()


def test_block_index_is_unique_per_document_and_allows_legacy_nulls(
    processing_context, migrated_connection
):
    """同一文档有序块不能重号，不同文档可同序号，多个历史空序号不被伪造或拒绝。"""

    first = upload_sample(processing_context, "顺序一.txt", b"first index fixture")
    second = upload_sample(processing_context, "顺序二.txt", b"second index fixture")
    connection, _, _ = migrated_connection

    def block(document, block_index):
        """构造数据库约束样本，保持其他字段合法并显式指定空或正数序号。"""

        return DocumentBlock(
            doc_id=UUID(document["doc_id"]),
            block_index=block_index,
            page=1,
            block_type="text",
            text="顺序约束正文",
            source="native_parser",
        )

    with Session(connection, join_transaction_mode="create_savepoint") as session:
        session.add_all(
            [block(first, 1), block(first, None), block(first, None), block(second, 1)]
        )
        session.flush()
        with pytest.raises(IntegrityError), session.begin_nested():
            session.add(block(first, 1))
            session.flush()
        assert session.scalar(select(func.count()).select_from(DocumentBlock)) == 4


def test_block_order_migration_keeps_legacy_data_without_inventing_order(
    processing_context, migrated_connection
):
    """0002 和 0003 往返保留旧块、文档及任务，历史顺序保持未知而不按 UUID 编造。"""

    document = upload_sample(
        processing_context, "历史文档.txt", b"legacy block fixture"
    )
    connection, config, schema = migrated_connection
    with connection.begin_nested():
        command.downgrade(config, "0002_document_tasks")
        assert "block_index" not in {
            column["name"]
            for column in inspect(connection).get_columns(
                "document_blocks", schema=schema
            )
        }
        old_ids = [uuid4(), uuid4()]
        for position, block_id in enumerate(old_ids):
            connection.execute(
                text(
                    "INSERT INTO document_blocks "
                    "(block_id, doc_id, page, block_type, text, source) "
                    "VALUES (:block_id, :doc_id, 1, 'text', :content, 'native_parser')"
                ),
                {
                    "block_id": block_id,
                    "doc_id": UUID(document["doc_id"]),
                    "content": f"旧正文-{position}",
                },
            )
        command.upgrade(config, "head")
        assert connection.execute(
            text("SELECT block_id, block_index FROM document_blocks ORDER BY text")
        ).all() == [(block_id, None) for block_id in old_ids]
        assert connection.scalar(
            text("SELECT processing_task_id FROM documents")
        ) == UUID(document["task_id"])
        command.check(config)
        command.downgrade(config, "0002_document_tasks")
        assert connection.scalar(text("SELECT count(*) FROM document_blocks")) == 2
        assert connection.execute(
            text("SELECT text FROM document_blocks ORDER BY text")
        ).scalars().all() == [
            "旧正文-0",
            "旧正文-1",
        ]
        command.upgrade(config, "head")
        assert (
            connection.scalar(
                text("SELECT count(*) FROM document_blocks WHERE block_index IS NULL")
            )
            == 2
        )
        command.check(config)

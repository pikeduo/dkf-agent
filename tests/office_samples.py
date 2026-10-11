"""离线 Office 样本；旧 Office 仅构造可辨识 CFB 容器，不能冒充真实云端验收原件。"""

import struct
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from openpyxl import Workbook


def ole_bytes(file_type="doc", *, marker=None) -> bytes:
    """生成规范 FAT/目录/流的最小 CFB，供类型识别测试，不声称 Office 正文可完整渲染。"""

    marker = (
        marker
        or {"doc": "WordDocument", "xls": "Workbook", "ppt": "PowerPoint Document"}[
            file_type
        ]
    )
    free, end = 0xFFFFFFFF, 0xFFFFFFFE
    header = bytearray(512)
    header[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<HHHHH", header, 24, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<IIIIIIIII", header, 40, 0, 1, 1, 0, 4096, end, 0, end, 0)
    struct.pack_into("<109I", header, 76, 0, *([free] * 108))
    fat = [0xFFFFFFFD, end] + list(range(3, 10)) + [end] + [free] * 118
    directory = bytearray(512)

    def entry(offset, name, kind, child=free, start=end, size=0):
        """填充一个目录项，使用根节点与单个黑色流节点构造可读取目录树。"""

        encoded = (name + "\0").encode("utf-16-le")
        directory[offset : offset + len(encoded)] = encoded
        struct.pack_into(
            "<HBBIII", directory, offset + 64, len(encoded), kind, 1, free, free, child
        )
        struct.pack_into("<IQ", directory, offset + 116, start, size)

    entry(0, "Root Entry", 5, child=1)
    entry(128, marker, 2, start=2, size=4096)
    payload = {
        "doc": b"\xec\xa5",
        "xls": b"\x09\x08\x08\x00",
        "ppt": b"\x0f\x00\xe8\x03",
    }[file_type]
    return (
        bytes(header)
        + struct.pack("<128I", *fat)
        + bytes(directory)
        + payload.ljust(4096, b"\0")
    )


def office_bytes(file_type) -> bytes:
    """生成真实 DOCX/XLSX 或最小 PPTX 主部件；PPTX 仅用于离线容器及 Mock 协议验收。"""

    if file_type in {"doc", "xls", "ppt"}:
        return ole_bytes(file_type)
    buffer = BytesIO()
    if file_type == "docx":
        document = Document()
        document.add_paragraph("Office 样本正文")
        document.save(buffer)
    elif file_type == "xlsx":
        workbook = Workbook()
        workbook.active.title = "中文数据"
        workbook.active.append(["项目", "数值", "占比", "公式"])
        workbook.active.append(["销售", 123, 0.25, "=B2*2"])
        workbook.active["C2"].number_format = "0%"
        workbook.save(buffer)
        workbook.close()
    elif file_type == "pptx":
        with ZipFile(buffer, "w") as archive:
            archive.writestr(
                "[Content_Types].xml",
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
            )
            archive.writestr(
                "ppt/presentation.xml",
                '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>',
            )
    else:
        raise ValueError("未定义的 Office 样本类型")
    return buffer.getvalue()

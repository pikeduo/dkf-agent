"""通过真实 DOCX 文件验证原生解析及统一结构，不需要数据库或 Redis。"""

from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.parsers import docx as docx_parser
from app.parsers.common import NoTextError, ParserError
from app.parsers.docx import DocxParser


def _parse(path: Path, doc_id: UUID | None = None):
    """通过公共校验入口读取临时 DOCX，保留显式原始文件名。"""

    return DocxParser().parse(
        doc_id=doc_id or uuid4(),
        file_path=path,
        file_name="原始文档.docx",
        file_type="docx",
    )


def _set_outline(element, value: str) -> None:
    """在测试段落或样式中写入大纲级别，模拟 Word 自定义标题。"""

    properties = element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), value)
    properties.append(outline)


def _rewrite_part(path: Path, part_name: str, replacement: bytes) -> None:
    """改写测试 DOCX 的一个 ZIP 部件，构造损坏 XML 或异常包类型。"""

    buffer = BytesIO()
    with ZipFile(path) as original, ZipFile(buffer, "w") as rewritten:
        for member in original.infolist():
            rewritten.writestr(
                member,
                replacement
                if member.filename == part_name
                else original.read(member.filename),
            )
    path.write_bytes(buffer.getvalue())


def test_docx_preserves_title_paragraph_table_and_section_order(tmp_path):
    """验证文档标题、段落、表格和章节层级按正文原始顺序输出。"""

    path = tmp_path / "stored-uuid.bin"
    document = Document()
    document.add_heading("管理报告", 0)
    document.add_paragraph("引言文字")
    document.add_heading("第一章", 1)
    document.add_paragraph("第一段\n第二行")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "收入"
    table.cell(1, 1).text = "100"
    document.add_heading("第一节", 2)
    document.add_paragraph("章节内容")
    document.add_heading("第二章", 1)
    document.add_paragraph("新章节内容")
    document.save(path)

    parsed = _parse(path)

    assert parsed.title == "管理报告"
    assert parsed.file_name == "原始文档.docx"
    assert parsed.file_type == "docx"
    assert [block.block_type for block in parsed.blocks] == [
        "title",
        "text",
        "title",
        "text",
        "table",
        "title",
        "text",
        "title",
        "text",
    ]
    assert parsed.blocks[3].text == "第一段\n第二行"
    assert parsed.blocks[4].text == "指标\t金额\n收入\t100"
    assert parsed.blocks[4].section == "管理报告 / 第一章"
    assert parsed.blocks[6].section == "管理报告 / 第一章 / 第一节"
    assert parsed.blocks[8].section == "管理报告 / 第二章"
    assert all(block.page == 1 for block in parsed.blocks)
    assert all(block.bbox is None for block in parsed.blocks)
    assert all(block.source == "native_parser" for block in parsed.blocks)
    assert all(block.confidence == 1.0 for block in parsed.blocks)


@pytest.mark.parametrize("level", range(1, 10))
def test_docx_recognizes_all_standard_heading_levels(tmp_path, level):
    """验证标准 Heading 1 到 Heading 9 都生成标题块。"""

    path = tmp_path / "headings.docx"
    document = Document()
    document.add_heading("层级标题", level)
    document.add_paragraph("正文")
    document.save(path)

    parsed = _parse(path)

    assert parsed.title == "层级标题"
    assert parsed.blocks[0].block_type == "title"
    assert parsed.blocks[1].section == "层级标题"


@pytest.mark.parametrize("style_name", ["标题", "标题 1", "标题2", "Heading 3"])
def test_docx_recognizes_localized_heading_styles(tmp_path, style_name):
    """验证中文标题样式与标准英文样式通过名称识别为标题。"""

    path = tmp_path / "localized.docx"
    document = Document()
    if style_name not in document.styles:
        document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("中文标题", style=style_name)
    document.add_paragraph("中文内容")
    document.save(path)

    parsed = _parse(path)

    assert parsed.blocks[0].block_type == "title"
    assert parsed.blocks[1].section == "中文标题"


def test_docx_recognizes_custom_inherited_heading(tmp_path):
    """验证继承标准标题的自定义样式仍保留章节信息。"""

    path = tmp_path / "custom.docx"
    document = Document()
    style = document.styles.add_style("报告章节", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Heading 2"]
    document.add_paragraph("自定义章节", style=style)
    document.add_paragraph("正文")
    document.save(path)

    assert _parse(path).blocks[0].block_type == "title"


@pytest.mark.parametrize("target", ["paragraph", "style"])
def test_docx_recognizes_explicit_outline_level(tmp_path, target):
    """验证段落或自定义样式中的大纲级别可以提供标题语义。"""

    path = tmp_path / "outline.docx"
    document = Document()
    paragraph = document.add_paragraph("大纲标题")
    if target == "paragraph":
        _set_outline(paragraph._p, "2")
    else:
        style = document.styles.add_style("大纲样式", WD_STYLE_TYPE.PARAGRAPH)
        _set_outline(style.element, "2")
        paragraph.style = style
    document.save(path)

    assert _parse(path).blocks[0].block_type == "title"


def test_docx_explicit_body_outline_overrides_heading_style(tmp_path):
    """验证显式正文大纲级别不会被继承标题样式误判。"""

    path = tmp_path / "body.docx"
    document = Document()
    paragraph = document.add_heading("这是正文", 1)
    _set_outline(paragraph._p, "9")
    document.save(path)

    parsed = _parse(path)

    assert parsed.title == ""
    assert parsed.blocks[0].block_type == "text"


def test_docx_nested_table_and_merged_cells_keep_text(tmp_path):
    """验证嵌套表格正文不丢失，水平合并单元格不重复输出。"""

    path = tmp_path / "tables.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并标题"
    table.cell(1, 0).text = "正文"
    nested = table.cell(1, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "内部A"
    nested.cell(0, 1).text = "内部B"
    document.save(path)

    parsed = _parse(path)

    assert len(parsed.blocks) == 1
    assert parsed.blocks[0].block_type == "table"
    assert parsed.blocks[0].text.count("合并标题") == 1
    assert "正文\t内部A\t内部B" in parsed.blocks[0].text


def test_docx_skips_empty_content_without_fake_physical_page(tmp_path):
    """验证空段落、分页符和空表格不会生成虚假块或物理页码。"""

    path = tmp_path / "blank-parts.docx"
    document = Document()
    document.add_paragraph("  \t ")
    document.add_table(rows=1, cols=2)
    document.add_page_break()
    document.add_paragraph("唯一正文")
    document.save(path)

    parsed = _parse(path)

    assert len(parsed.blocks) == 1
    assert parsed.blocks[0].text == "唯一正文"
    assert parsed.blocks[0].page == 1


def test_docx_preserves_body_and_cell_whitespace(tmp_path):
    """验证正文与表格单元格首尾空白及换行原样保留，仅标题元信息清理边界空白。"""

    path = tmp_path / "whitespace.docx"
    document = Document()
    document.add_heading("  报告标题  ", 0)
    document.add_paragraph("  正文首行\n第二行\t  ")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "  单元格A\n另一行  "
    table.cell(0, 1).text = "\t单元格B  "
    document.save(path)

    parsed = _parse(path)

    assert parsed.title == "报告标题"
    assert parsed.blocks[0].text == "  报告标题  "
    assert parsed.blocks[1].text == "  正文首行\n第二行\t  "
    assert parsed.blocks[1].section == "报告标题"
    assert parsed.blocks[2].text == "  单元格A\n另一行  \t\t单元格B  "


def test_docx_block_ids_are_repeatable_and_document_scoped(tmp_path):
    """验证重试解析不改变块标识，不同文档的相同正文不会共享块标识。"""

    path = tmp_path / "repeat.docx"
    document = Document()
    document.add_paragraph("相同内容")
    document.add_paragraph("相同内容")
    document.save(path)
    doc_id = uuid4()

    first = _parse(path, doc_id)
    second = _parse(path, doc_id)
    different = _parse(path)

    assert first.model_dump() == second.model_dump()
    assert len({block.block_id for block in first.blocks}) == 2
    assert first.blocks[0].block_id != different.blocks[0].block_id


def test_docx_empty_document_reports_no_text(tmp_path):
    """验证合法空 DOCX 不被误报为解析成功。"""

    path = tmp_path / "empty.docx"
    Document().save(path)

    with pytest.raises(NoTextError, match="未提取到原生正文文字"):
        _parse(path)


@pytest.mark.parametrize("payload", [b"", b"not-a-docx", b"PK\x03\x04broken"])
def test_docx_invalid_package_has_safe_error(tmp_path, payload):
    """验证空文件与损坏 ZIP 只返回安全中文错误而不泄漏存储路径。"""

    path = tmp_path / "private-file.docx"
    path.write_bytes(payload)

    with pytest.raises(ParserError) as error:
        _parse(path)

    assert "DOCX 文件损坏" in str(error.value)
    assert str(path) not in str(error.value)


def test_docx_missing_package_part_has_safe_error(tmp_path):
    """验证缺少 OPC 内容类型部件的 ZIP 包被识别为损坏文档。"""

    path = tmp_path / "missing-part.docx"
    with ZipFile(path, "w") as package:
        package.writestr("word/document.xml", "<x/>")

    with pytest.raises(ParserError, match="DOCX 文件损坏"):
        _parse(path)


def test_docx_broken_xml_has_safe_error(tmp_path):
    """验证损坏的正文 XML 被转换成安全解析错误。"""

    path = tmp_path / "broken-xml.docx"
    Document().save(path)
    _rewrite_part(path, "word/document.xml", b"<broken>")

    with pytest.raises(ParserError, match="DOCX 文件损坏"):
        _parse(path)


def test_docx_missing_source_keeps_os_error(tmp_path):
    """验证缺少原件时保留 FileNotFoundError，供任务层区分错误与重试。"""

    with pytest.raises(FileNotFoundError):
        _parse(tmp_path / "missing.docx")


def test_docx_loader_os_error_is_not_wrapped(tmp_path, monkeypatch):
    """验证底层临时 I/O 错误不会被错误归类为永久格式损坏。"""

    path = tmp_path / "access.docx"
    Document().save(path)

    def failed_loader(source):
        """模拟依赖库读取原件时遇到临时系统 I/O 错误。"""

        raise OSError("temporary read error")

    monkeypatch.setattr(docx_parser, "load_document", failed_loader)

    with pytest.raises(OSError, match="temporary read error"):
        _parse(path)


def test_docx_wrong_file_type_is_rejected_before_open(tmp_path):
    """验证类型由调用方显式提供，错误类型不会触发磁盘访问。"""

    with pytest.raises(ParserError, match="只支持 docx"):
        DocxParser().parse(
            doc_id=uuid4(),
            file_path=tmp_path / "missing.docx",
            file_name="错误类型.txt",
            file_type="txt",
        )

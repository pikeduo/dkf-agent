"""CSV/XLSX Native 契约、字段保真、错误处理及稳定标识的离线测试。"""

import codecs
import csv
from io import BytesIO, StringIO
from uuid import uuid4

import pytest
from csv_samples import invalid_encoding_bytes, regenerate_manual_fixture
from office_samples import office_bytes
from openpyxl import Workbook

from app.parsers.common import ParserError
from app.parsers.csv import CsvParser
from app.parsers.registry import get_parser
from app.parsers.text import read_text
from app.parsers.xlsx import XlsxParser


def parsed_file(tmp_path, content, file_type, *, doc_id=None):
    """保存临时原件后走公共 parse 入口，验证与生产一致的归属及字段校验。"""

    path = tmp_path / f"样本.{file_type}"
    path.write_bytes(content)
    return get_parser(file_type).parse(
        doc_id=doc_id or uuid4(),
        file_path=path,
        file_name=path.name,
        file_type=file_type,
    )


def rows_of(block):
    """按 TSV 的引号规则读回 Block 表格，避免使用 split 掩盖字段内换行或制表符。"""

    return list(
        csv.reader(StringIO(block.text, newline=""), delimiter="\t", strict=True)
    )


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "gb18030"])
def test_csv_preserves_fields_and_stable_ids(tmp_path, encoding):
    """三类编码保留中文、空值、引号、逗号、换行、数字及百分比，重试不改变 ID。"""

    content = '名称,备注,数值,占比\r\n中文,"逗号,与""引号""",123,25%\r\n空值,,0,0%\r\n换行,"第一行\n第二行",1,100%\r\n'.encode(
        encoding
    )
    doc_id = uuid4()
    parsed = parsed_file(tmp_path, content, "csv", doc_id=doc_id)
    assert rows_of(parsed.blocks[0]) == [
        ["名称", "备注", "数值", "占比"],
        ["中文", '逗号,与"引号"', "123", "25%"],
        ["空值", "", "0", "0%"],
        ["换行", "第一行\n第二行", "1", "100%"],
    ]
    assert parsed == parsed_file(tmp_path, content, "csv", doc_id=doc_id)
    block = parsed.blocks[0]
    assert (block.page, block.block_type, block.bbox, block.source) == (
        1,
        "table",
        None,
        "native_parser",
    )


@pytest.mark.parametrize(
    "content",
    [
        b"",
        b"\xef\xbb\xbf",
        b",,\n,,",
        invalid_encoding_bytes(),
        b"a,b\n1",
        b'a,b\n"unclosed,2',
        b'a,b\n"x"bad,2',
        b"a,\0b",
    ],
)
def test_csv_invalid_inputs_fail_safely(tmp_path, content):
    """空 CSV、非法编码、控制字符和畸形结构必须明确失败，不返回空成功结果。"""

    with pytest.raises(ParserError):
        parsed_file(tmp_path, content, "csv")


@pytest.mark.parametrize("encoding", ["utf-8", "gb18030"])
def test_invalid_encoding_fixture_is_really_undecodable(encoding):
    """先验证负向字节本身无 BOM 且严格解码失败，防止合法编码误当负向样本。"""

    content = invalid_encoding_bytes()
    assert content == b"\x81\x30\x81"
    assert not content.startswith(
        (
            codecs.BOM_UTF8,
            codecs.BOM_UTF16_LE,
            codecs.BOM_UTF16_BE,
            codecs.BOM_UTF32_LE,
            codecs.BOM_UTF32_BE,
        )
    )
    with pytest.raises(UnicodeDecodeError):
        content.decode(encoding, errors="strict")


@pytest.mark.parametrize("entry_point", ["read_text", "CsvParser"])
def test_invalid_encoding_errors_are_safe(tmp_path, entry_point):
    """真实解码失败返回固定安全 ParserError，不泄漏输入字节、文件名或本机路径。"""

    content = invalid_encoding_bytes()
    path = tmp_path / "私有路径样本.csv"
    path.write_bytes(content)
    with pytest.raises(ParserError) as error:
        if entry_point == "read_text":
            read_text(path)
        else:
            CsvParser().parse(
                doc_id=uuid4(), file_path=path, file_name=path.name, file_type="csv"
            )
    message = str(error.value)
    assert message == "文本编码无法识别，请转换为 UTF-8 后重新上传"
    assert repr(content) not in message and "\\x81" not in message
    assert str(path) not in message and path.name not in message
    assert error.value.__cause__ is None


def test_valid_utf16_unreadable_text_is_not_encoding_failure(tmp_path):
    """合法 UTF-16 即使语义不可读也不是解码失败；第一版不按具体字符串猜测乱码。"""

    content = codecs.BOM_UTF16_LE + "\u8180\u8382".encode("utf-16-le")
    parsed = parsed_file(tmp_path, content, "csv")
    assert rows_of(parsed.blocks[0]) == [[content.decode("utf-16", errors="strict")]]


@pytest.mark.parametrize("prefix", ["", "samples/"])
def test_manual_fixture_regeneration_preserves_other_zip_entries(tmp_path, prefix):
    """更新二进制负向样本与说明时，ZIP 内其他验收素材、压缩方式和注释保持不变。"""

    from zipfile import ZIP_DEFLATED, ZipFile

    sample_dir = tmp_path / "samples"
    sample_dir.mkdir()
    name = "04_negative_invalid_encoding.csv"
    old_content = codecs.BOM_UTF16_LE + "\u8180\u8382".encode("utf-16-le")
    (sample_dir / name).write_bytes(old_content)
    (sample_dir / "README_TEST_EXPECTATIONS.md").write_text(
        "修正后的说明", encoding="utf-8"
    )
    archive_path = tmp_path / "samples.zip"
    with ZipFile(archive_path, "w", compression=ZIP_DEFLATED) as archive:
        archive.comment = b"preserve-comment"
        archive.writestr(prefix + name, old_content)
        archive.writestr(prefix + "README_TEST_EXPECTATIONS.md", b"old expectations")
        archive.writestr(prefix + "other.xlsx", b"OTHER_SAMPLE_BYTES")
    regenerate_manual_fixture(sample_dir, archive_path)
    assert (sample_dir / name).read_bytes() == invalid_encoding_bytes()
    with ZipFile(archive_path) as archive:
        assert archive.read(prefix + name) == invalid_encoding_bytes()
        assert (
            archive.read(prefix + "README_TEST_EXPECTATIONS.md").decode()
            == "修正后的说明"
        )
        assert archive.read(prefix + "other.xlsx") == b"OTHER_SAMPLE_BYTES"
        assert archive.getinfo(prefix + "other.xlsx").compress_type == ZIP_DEFLATED
        assert archive.comment == b"preserve-comment"


def test_xlsx_cells_formulas_and_percentages(tmp_path):
    """按原值提取中文和数字，保留普通百分比显示及未执行的公式文本。"""

    doc_id = uuid4()
    parsed = parsed_file(tmp_path, office_bytes("xlsx"), "xlsx", doc_id=doc_id)
    assert rows_of(parsed.blocks[0]) == [
        ["项目", "数值", "占比", "公式"],
        ["销售", "123", "25%", "=B2*2"],
    ]
    assert parsed.blocks[0].section == "中文数据"
    assert parsed == parsed_file(tmp_path, office_bytes("xlsx"), "xlsx", doc_id=doc_id)
    assert parsed.blocks[0].bbox is None and parsed.blocks[0].page == 1


def test_xlsx_sheet_order_empty_and_cell_positions(tmp_path):
    """保留工作表原序及内部空行空列，跳过空 Sheet，不把 Sheet 当物理页码。"""

    workbook = Workbook()
    workbook.active.title = "第一表"
    workbook.active["B2"] = "右下"
    workbook.create_sheet("空表")
    workbook.create_sheet("第三表").append(["中文", 3, "=1+2"])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    parsed = parsed_file(tmp_path, buffer.getvalue(), "xlsx")
    assert [block.section for block in parsed.blocks] == ["第一表", "第三表"]
    assert rows_of(parsed.blocks[0]) == [["", ""], ["", "右下"]]
    assert rows_of(parsed.blocks[1]) == [["中文", "3", "=1+2"]]
    assert all(
        block.page == 1 and block.bbox is None and block.block_type == "table"
        for block in parsed.blocks
    )


def test_xlsx_empty_workbook_fails(tmp_path):
    """全空工作簿不能作为完成解析的证据。"""

    workbook = Workbook()
    workbook.create_sheet("第二空表")
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    with pytest.raises(ParserError, match="均无"):
        parsed_file(tmp_path, buffer.getvalue(), "xlsx")


@pytest.mark.parametrize("file_type", ["doc", "xls", "ppt", "pptx", "png"])
def test_unsupported_native_formats_require_explicit_mineru(file_type):
    """旧 Office、演示文稿及图片不得注册 Native 或隐式启动云解析。"""

    with pytest.raises(ParserError, match="MinerU"):
        get_parser(file_type)


@pytest.mark.parametrize(
    "parser,file_type", [(CsvParser(), "txt"), (XlsxParser(), "xls")]
)
def test_tabular_parser_rejects_wrong_type(tmp_path, parser, file_type):
    """直接调用 Parser 也不能绕过格式边界。"""

    with pytest.raises(ParserError):
        parser.parse(
            doc_id=uuid4(),
            file_path=tmp_path / "missing",
            file_name="wrong",
            file_type=file_type,
        )

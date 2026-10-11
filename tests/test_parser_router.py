"""整份路由、结构预检和完整性保护；不访问真实 Broker、数据库或云 API。"""

from uuid import uuid4

import pytest
from mineru_samples import png_bytes
from office_samples import office_bytes
from pydantic import ValidationError
from router_samples import docx_structure, office_with_part, router_pdf

from app.core.config import Settings
from app.parsers.common import ParserError, make_block
from app.parsers.registry import get_parser
from app.parsers.routing.contracts import Inspection, RouteDecision, StructuralSignals
from app.parsers.routing.guard import CompletenessError, guard_result
from app.parsers.routing.inspector import inspect_document
from app.parsers.routing.router import decide_route, fallback_route
from app.schemas.parsed_document import ParsedDocument


def inspect_sample(tmp_path, kind, content, **configuration):
    """写入独立临时原件，以无密钥配置执行真实预检和纯路由。"""

    path = tmp_path / ("任意名称." + kind)
    path.write_bytes(content)
    inspected = inspect_document(path, kind, Settings(_env_file=None, **configuration))
    return path, inspected, decide_route(kind, "auto", inspected)


@pytest.mark.parametrize(
    "kind,content",
    [
        ("txt", b"plain text"),
        ("csv", b"name,value\na,1"),
        ("md", b"# title\n\nplain text"),
        ("docx", office_bytes("docx")),
        ("xlsx", office_bytes("xlsx")),
        ("pdf", router_pdf()),
    ],
    ids=["txt", "csv", "md", "docx", "xlsx", "pdf"],
)
def test_simple_native_matrix_and_guard(tmp_path, kind, content):
    """简单六种 Native 文档路由稳定，预检后真实 Parser 结果通过落库前检查。"""

    path, inspected, decision = inspect_sample(tmp_path, kind, content)
    assert decision.route_strategy == "native"
    assert decision.reason_codes == ("NATIVE_COVERAGE_SUPPORTED",)
    assert decide_route(kind, "auto", inspected) == decision
    parsed = get_parser(kind).parse(
        doc_id=uuid4(), file_name=path.name, file_type=kind, file_path=path
    )
    guard_result(parsed, decision)
    if kind == "xlsx":
        assert (
            decision.signals.sheet_count == decision.signals.nonempty_sheet_count == 1
        )
        assert decision.signals.formula_count == 1 and "=B2*2" in parsed.blocks[0].text


@pytest.mark.parametrize("kind", ["jpg", "jpeg", "png", "doc", "xls", "ppt", "pptx"])
def test_visual_legacy_matrix(kind):
    """图片、旧 Office 与演示文稿统一整份 MinerU，不注册二进制 Native Parser。"""

    decision = decide_route(kind, "auto", Inspection())
    assert decision.route_strategy == "mineru_whole"
    assert decision.reason_codes == ("VISUAL_OR_LEGACY_FORMAT",)


@pytest.mark.parametrize(
    "content,code",
    [
        (docx_structure("<w:p><w:r><w:drawing/></w:r></w:p>"), "DOCX_HAS_DRAWING"),
        (
            docx_structure(
                '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"/>'
            ),
            "DOCX_HAS_FORMULA",
        ),
        (docx_structure("<w:txbxContent><w:p/></w:txbxContent>"), "DOCX_HAS_TEXTBOX"),
        (
            office_with_part("docx", "word/media/asset.png", png_bytes()),
            "DOCX_HAS_EMBEDDED_IMAGE",
        ),
        (
            office_with_part(
                "docx",
                "word/header1.xml",
                '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>header</w:t></w:r></w:p></w:hdr>',
            ),
            "DOCX_HAS_AUXILIARY_CONTENT",
        ),
        (
            docx_structure("<w:sdt><w:sdtContent><w:p/></w:sdtContent></w:sdt>"),
            "DOCX_HAS_UNCOVERED_STRUCTURE",
        ),
    ],
    ids=["drawing", "formula", "textbox", "media", "header", "content_control"],
)
def test_docx_uncovered_structure_whole(tmp_path, content, code):
    """正文很多也不能忽略 DOCX 图片、公式、文本框和未覆盖的辅助结构。"""

    _, _, decision = inspect_sample(tmp_path, "docx", content)
    assert decision.route_strategy == "mineru_whole" and code in decision.reason_codes


@pytest.mark.parametrize(
    "name,code",
    [
        ("xl/media/image.png", "XLSX_HAS_EMBEDDED_IMAGE"),
        ("xl/charts/chart1.xml", "XLSX_HAS_CHART"),
        ("xl/drawings/drawing1.xml", "XLSX_HAS_DRAWING"),
        ("xl/embeddings/object.bin", "XLSX_HAS_EMBEDDED_OBJECT"),
        ("xl/drawings/shape.vml", "XLSX_HAS_VML"),
        ("xl/externalLinks/externalLink1.xml", "XLSX_HAS_EXTERNAL_LINK"),
    ],
)
def test_complex_xlsx_hybrid_or_missing_asset_failure(tmp_path, name, code):
    """复杂 XLSX 优先 Hybrid 预检；外部工作簿未上传则拒绝，不让普通用户选模式。"""

    content = office_with_part("xlsx", name, "<root/>")
    _, inspected, decision = inspect_sample(tmp_path, "xlsx", content)
    external = code == "XLSX_HAS_EXTERNAL_LINK"
    assert decision.route_strategy == ("failed" if external else "hybrid_bundle")
    assert code in decision.reason_codes
    assert decide_route("xlsx", "mineru", inspected).route_strategy == ("failed" if external else "mineru_whole")
    assert decide_route("xlsx", "native", inspected).route_strategy == "native"


@pytest.mark.parametrize(
    "kind,code",
    [
        ("scan", "PDF_SCAN_OR_NON_TEXT_PAGE"),
        ("same_page", "PDF_TEXT_IMAGE_MIXED"),
        ("mixed_pages", "PDF_SCAN_OR_NON_TEXT_PAGE"),
        ("columns", "PDF_MULTICOLUMN_RISK"),
        ("drawing", "PDF_COMPLEX_DRAWING"),
    ],
)
def test_pdf_page_level_risks(tmp_path, kind, code):
    """扫描、跨页/同页混合、多栏和大绘图走整份 MinerU，保留物理页摘要。"""

    _, _, decision = inspect_sample(tmp_path, "pdf", router_pdf(kind))
    assert decision.route_strategy == "mineru_whole" and code in decision.reason_codes
    assert len(decision.signals.pages) == (3 if kind == "mixed_pages" else 1)
    assert [page.page for page in decision.signals.pages] == list(
        range(1, len(decision.signals.pages) + 1)
    )


def test_pdf_small_logo_threshold_and_limits(tmp_path):
    """小 logo 不单独触发云端；阈值可调，资源超限明确失败 而非部分预检。"""

    _, _, decision = inspect_sample(tmp_path, "pdf", router_pdf("logo"))
    assert decision.route_strategy == "native"
    _, _, sensitive = inspect_sample(
        tmp_path, "pdf", router_pdf("logo"), router_pdf_image_area_ratio=0.0001
    )
    assert sensitive.route_strategy == "mineru_whole"
    _, _, limited = inspect_sample(
        tmp_path, "pdf", router_pdf("mixed_pages"), router_pdf_max_pages=1
    )
    assert limited.route_strategy == "failed" and limited.reason_codes == (
        "PREFLIGHT_LIMIT",
    )


@pytest.mark.parametrize(
    "content,rejected",
    [
        ("# 文本\n\n![图](images/a.png)", True),
        ("![图](https://example.invalid/asset.png)", True),
        ("![引用][asset]", True),
        ('<img src="asset.png">', True),
        ("```md\n![示例](not-an-asset)\n```\n正文", False),
        ("`![示例](not-an-asset)` 正文", False),
        (r"\![普通文字]", False),
    ],
)
def test_markdown_assets_without_fetch(tmp_path, content, rejected):
    """图片引用明确失败，围栏示例不误判；不把 MD 送往 MinerU 或抓取远程附件。"""

    _, _, decision = inspect_sample(tmp_path, "md", content.encode())
    assert decision.route_strategy == ("failed" if rejected else "native")
    assert (
        "https://" not in decision.model_dump_json()
        and "asset.png" not in decision.model_dump_json()
    )


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_missing_external_attachment_is_terminal_even_for_cloud_override(kind):
    """单文件无法覆盖外部附件时 auto/显式云均拒绝，不静默遗漏或抓取 URL。"""

    inspection = Inspection(risks=("EXTERNAL_ATTACHMENT_MISSING",))
    for mode in ("auto", "mineru"):
        decision = decide_route(kind, mode, inspection)
        assert decision.route_strategy == "failed"
        assert decision.reason_codes == ("EXTERNAL_ATTACHMENT_MISSING",)


def test_office_relationship_stub_distinguishes_attachment_and_hyperlink(monkeypatch):
    """只用关系 XML 元数据检查外链/缺部件，不创建或打开 Office 文件。"""

    from types import SimpleNamespace
    from lxml import etree
    from app.parsers.routing import inspector

    relationship = {"Id": "rId1", "Target": "https://private.invalid/asset", "TargetMode": "External",
                    "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"}

    class PackageStub:
        """仅提供关系摘要，不包含压缩包或真实文档。"""

        def namelist(self):
            """返回一个关系文件和一个本地目标名。"""
            return ["xl/_rels/workbook.xml.rels", "xl/worksheets/s1.xml"]

        def getinfo(self, name):
            """提供固定安全大小，不访问文件。"""
            return SimpleNamespace(file_size=1)

    def xml_stub(archive, name, limit):
        """直接构建内存关系节点，不解析 OOXML 文件。"""
        root = etree.Element("Relationships")
        etree.SubElement(root, "Relationship", **relationship)
        return root

    monkeypatch.setattr(inspector, "xml_part", xml_stub)
    archive = PackageStub()
    assert inspector.missing_external_assets(archive, "xl", 1024)
    relationship["Type"] = relationship["Type"].rsplit("/", 1)[0] + "/hyperlink"
    assert not inspector.missing_external_assets(archive, "xl", 1024)
    relationship.update(TargetMode="Internal", Target="worksheets/s1.xml", Type="worksheet")
    assert not inspector.missing_external_assets(archive, "xl", 1024)
    relationship["Target"] = "worksheets/missing-private.xml"
    with pytest.raises(ParserError) as caught:
        inspector.missing_external_assets(archive, "xl", 1024)
    assert "missing-private" not in str(caught.value) and "https://" not in str(caught.value)


def test_explicit_overrides_and_fallback_once():
    """显式选择保留原因但不绕过契约；auto Native 只允许一次明确完整性回退。"""

    risk = Inspection(risks=("DOCX_HAS_DRAWING",))
    assert decide_route("docx", "native", risk).reason_codes == (
        "EXPLICIT_NATIVE",
        "DOCX_HAS_DRAWING",
    )
    for kind in ("txt", "csv", "md"):
        assert decide_route(kind, "mineru", Inspection()).route_strategy == "failed"
    original = decide_route("pdf", "auto", Inspection())
    changed = fallback_route(original)
    assert changed.fallback_count == 1 and changed.route_strategy == "mineru_whole"
    assert changed.reason_codes[-1] == "NATIVE_COMPLETENESS_FALLBACK"
    for rejected in (
        changed,
        decide_route("pdf", "native", Inspection()),
        decide_route("csv", "auto", Inspection()),
    ):
        with pytest.raises(ValueError):
            fallback_route(rejected)


def test_guard_rejects_missing_pages_without_reordering():
    """非空物理页不能静默漏掉；检查不跨页重排、不拆表格公式或随机改变 ID。"""

    decision = decide_route(
        "pdf",
        "auto",
        Inspection(
            signals=StructuralSignals(
                page_count=2,
                pages=(
                    {
                        "page": 1,
                        "text_chars": 10,
                        "image_count": 0,
                        "image_area_ratio": 0,
                        "drawing_count": 0,
                        "annotation_count": 0,
                        "required": True,
                    },
                    {
                        "page": 2,
                        "text_chars": 10,
                        "image_count": 0,
                        "image_area_ratio": 0,
                        "drawing_count": 0,
                        "annotation_count": 0,
                        "required": True,
                    },
                ),
            )
        ),
    )
    identifier = uuid4()
    parsed = ParsedDocument(
        doc_id=identifier,
        file_name="文档.pdf",
        file_type="pdf",
        blocks=[make_block(identifier, 1, page=1, block_type="table", text="原生表格")],
    )
    with pytest.raises(CompletenessError):
        guard_result(parsed, decision)
    parsed.blocks.append(
        make_block(identifier, 2, page=2, block_type="formula", text="x + y")
    )
    before = parsed.model_dump_json()
    guard_result(parsed, decision)
    assert parsed.model_dump_json() == before


def test_office_guard_and_safe_contract():
    """Office 不接受假坐标，摘要契约拒绝 Token/URL/正文等额外字段。"""

    identifier = uuid4()
    block = make_block(
        identifier, 1, page=1, block_type="text", text="正文", bbox=[0, 0, 10, 10]
    ).model_copy(update={"source": "mineru_cloud"})
    parsed = ParsedDocument(
        doc_id=identifier, file_name="文档.docx", file_type="docx", blocks=[block]
    )
    decision = decide_route("docx", "mineru", Inspection())
    with pytest.raises(CompletenessError, match="坐标"):
        guard_result(parsed, decision)
    with pytest.raises(ValidationError):
        StructuralSignals(token="not allowed")
    with pytest.raises(ValidationError):
        RouteDecision(**(decision.model_dump() | {"signed_url": "https://invalid"}))
    with pytest.raises(ValidationError):
        RouteDecision(
            **(decision.model_dump() | {"reason_codes": ["https://secret-url"]})
        )


def test_damaged_or_entity_xml_safe_failure(tmp_path):
    """破损及实体 XML 安全失败，不透传本地路径/原文，不为坏文档盲目调用云 API。"""

    content = office_with_part(
        "docx",
        "word/document.xml",
        '<!DOCTYPE root [<!ENTITY x SYSTEM "file:///private">]><root>&x;</root>',
    )
    with pytest.raises(ParserError) as caught:
        inspect_sample(tmp_path, "docx", content)
    assert str(tmp_path) not in str(caught.value) and "private" not in str(caught.value)


def test_xml_inspection_limit_is_failure(tmp_path):
    """单部件超限安全失败，不能因结构无法全检而静默 Native 或盲目送云。"""

    content = office_with_part(
        "docx", "word/document.xml", "<root>" + "x" * (1024 * 1024) + "</root>"
    )
    _, _, decision = inspect_sample(tmp_path, "docx", content, router_xml_part_max_mb=1)
    assert decision.route_strategy == "failed" and decision.reason_codes == (
        "PREFLIGHT_LIMIT",
    )


def test_pdf_formula_annotation_and_empty_page_handling(tmp_path):
    """公式字体和批注有明确理由；真正空白页不虚构正文，也不压缩后续页码。"""

    import pymupdf

    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((40, 50), "Native main text with enough characters")
        page.insert_text((40, 80), "abc", fontname="symb")
        page.add_text_annot((40, 120), "annotation")
        document.new_page()
        content = document.tobytes()
    _, _, decision = inspect_sample(tmp_path, "pdf", content)
    assert decision.route_strategy == "mineru_whole"
    assert (
        "PDF_FORMULA_RISK" in decision.reason_codes
        and "PDF_HAS_ANNOTATION" in decision.reason_codes
    )
    assert [page.required for page in decision.signals.pages] == [True, False]


def test_xlsx_multisheet_empty_formula_and_guard(tmp_path):
    """多 Sheet 与空表顺序和公式保留；Guard 拒绝删除一个有效 Sheet 的结果。"""

    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "第一张"
    workbook.active.append(["中文", "=B2*2", 0.25])
    workbook.create_sheet("空白")
    workbook.create_sheet("只有空格").append(["  "])
    workbook.create_sheet("最后一张").append([123, "末页数据"])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    path, _, decision = inspect_sample(tmp_path, "xlsx", buffer.getvalue())
    assert (
        decision.route_strategy == "native"
        and decision.signals.nonempty_sheet_count == 2
    )
    parsed = get_parser("xlsx").parse(
        doc_id=uuid4(), file_path=path, file_type="xlsx", file_name=path.name
    )
    guard_result(parsed, decision)
    assert [block.section for block in parsed.blocks] == ["第一张", "最后一张"]
    assert all(block.page == 1 and block.bbox is None for block in parsed.blocks)
    parsed.blocks.pop()
    with pytest.raises(CompletenessError, match="Sheet"):
        guard_result(parsed, decision)
